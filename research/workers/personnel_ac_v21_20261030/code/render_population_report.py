"""Render the same evidence-backed report blocks to DOCX, HTML and Markdown.

This module is presentation only. It does not select teams, score geometry,
aggregate scientific outcomes, or change the evidence extraction ledger.
"""
from __future__ import annotations
import base64
import html
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'stage08_report'
STEM='HOHONET_v21_AC_population_report_ZH'
FONT='Noto Sans CJK SC'


def set_font(run,size=None,bold=None):
    run.font.name=FONT;run.font.color.rgb=RGBColor(0,0,0)
    if size is not None:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    rf=run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),FONT)


def cell_setup(cell,fill=None):
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    pr=cell._tc.get_or_add_tcPr()
    margin=OxmlElement('w:tcMar')
    for side,value in [('top',70),('bottom',70),('left',85),('right',85)]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');margin.append(e)
    pr.append(margin)
    if fill:
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),fill);pr.append(sh)


def docx(blocks):
    d=Document();sec=d.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.75);sec.bottom_margin=Cm(1.75)
    sec.left_margin=Cm(1.9);sec.right_margin=Cm(1.9)
    for name in ['Normal','Title','Heading 1','Heading 2','Caption','Header','Footer']:
        st=d.styles[name];st.font.name=FONT;st.font.color.rgb=RGBColor(0,0,0)
        rf=st._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),FONT)
    for border in d.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    st=d.styles['Normal'];st.font.size=Pt(10)
    st.paragraph_format.line_spacing=1.22;st.paragraph_format.space_after=Pt(5)
    for name,size in [('Title',22),('Heading 1',14),('Heading 2',11.5)]:
        st=d.styles[name];st.font.size=Pt(size);st.font.bold=True
        st.paragraph_format.keep_with_next=True
        st.paragraph_format.space_before=Pt(11);st.paragraph_format.space_after=Pt(6)
    hp=sec.header.paragraphs[0];set_font(hp.add_run('HOHONET v2.1  普通 Manual 人群模拟'),8)
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    set_font(fp.add_run('2026-10-29  ·  '),8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
    for b in blocks:
        if b['kind']=='h':
            p=d.add_paragraph(style='Title' if b['level']==0 else 'Heading '+str(b['level']))
            if b.get('page_break'):p.paragraph_format.page_break_before=True
            set_font(p.add_run(b['text']))
        elif b['kind']=='p':
            p=d.add_paragraph();set_font(p.add_run(b['text']),8.0 if b.get('small') else None)
            if b.get('small'):
                p.paragraph_format.line_spacing=1.06
                p.paragraph_format.space_after=Pt(2)
            if b.get('source'):set_font(p.add_run(' ['+b['source']+']'),8)
        elif b['kind']=='table':
            p=d.add_paragraph();p.paragraph_format.keep_with_next=True
            if b['title']=='可直接人工检查的具体问题':
                p.paragraph_format.page_break_before=True
            set_font(p.add_run(b['title']),10,True)
            t=d.add_table(rows=1,cols=len(b['headers']));t.autofit=False
            widths=b.get('widths') or [17.2/len(b['headers'])]*len(b['headers'])
            for col,width in zip(t.columns,widths):col.width=Cm(width)
            pr=t._tbl.tblPr;bounds=OxmlElement('w:tblBorders')
            for side in ['top','left','bottom','right','insideH','insideV']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bounds.append(e)
            pr.append(bounds)
            head=t.rows[0];repeat=OxmlElement('w:tblHeader');head._tr.get_or_add_trPr().append(repeat)
            for i,(c,txt) in enumerate(zip(head.cells,b['headers'])):
                c.width=Cm(widths[i]);cell_setup(c,'E7EDF3')
                p=c.paragraphs[0];p.paragraph_format.space_after=Pt(1)
                set_font(p.add_run(str(txt)),8.3,True)
            for row in b['rows']:
                rr=t.add_row();rr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                for i,(c,txt) in enumerate(zip(rr.cells,row)):
                    c.width=Cm(widths[i]);cell_setup(c)
                    p=c.paragraphs[0];p.paragraph_format.space_after=Pt(1)
                    p.paragraph_format.line_spacing=1.12
                    if i in b.get('center_columns',[]):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    set_font(p.add_run(str(txt)),b.get('font_size',8.3))
            if b.get('source'):
                p=d.add_paragraph();set_font(p.add_run('证据 ['+b['source']+']。完整精度与各指标分母保留在对应结果表。'),8)
        elif b['kind']=='image':
            path=ROOT/b['path'];p=d.add_paragraph();p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(path),width=Cm(b.get('width_cm',17.2)))
            p=d.add_paragraph(style='Caption');set_font(p.add_run(b['caption']),8.3)
    d.core_properties.title='HOHONET v2.1 直接AC与整类人群模拟研究'
    d.core_properties.subject='固定人数替换 整类替换 类别联盟 成员与背景敏感性'
    d.core_properties.author='HOHONET research'
    destination=OUT/(STEM+'.docx');d.save(destination)
    return destination


def web_and_markdown(blocks):
    css='body{font-family:system-ui,"Noto Sans CJK SC",sans-serif;color:#17212c;background:#fff;max-width:1100px;margin:40px auto;padding:0 26px;line-height:1.85}h1,h2,h3{color:#000;line-height:1.45}h1{font-size:31px}h2{font-size:23px;margin-top:40px}h3{font-size:18px}p{margin:13px 0}.tablewrap{overflow:auto}table{border-collapse:collapse;width:100%;font-size:14px;line-height:1.65}th,td{border:1px solid #d9d9d9;padding:9px 10px;vertical-align:middle}th{background:#e7edf3;text-align:left}figure{margin:25px 0}img{width:100%;height:auto}small,figcaption,.evidence{font-size:12px;color:#3f4c58;overflow-wrap:anywhere}@media print{body{max-width:none;font-size:10pt}h2,h3{break-after:avoid}tr{break-inside:avoid}}'
    result=['<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>HOHONET v2.1 直接AC与整类人群模拟研究</title><style>'+css+'</style><body>'];md=[]
    for b in blocks:
        k=b['kind'];src=' ['+b['source']+']' if b.get('source') else ''
        if k=='h':
            n=b['level']+1;result.append(f'<h{n}>'+html.escape(b['text'])+f'</h{n}>');md.append('#'*n+' '+b['text']+'\n')
        elif k=='p':
            result.append('<p>'+html.escape(b['text'])+('<small>'+html.escape(src)+'</small>' if src else '')+'</p>');md.append(b['text']+src+'\n')
        elif k=='table':
            result.append('<h3>'+html.escape(b['title'])+'</h3><div class="tablewrap"><table><thead><tr>'+''.join('<th>'+html.escape(str(x))+'</th>' for x in b['headers'])+'</tr></thead><tbody>')
            md.extend(['**'+b['title']+'**\n','| '+' | '.join(str(x).replace('|','／').replace('\n','<br>') for x in b['headers'])+' |','| '+' | '.join(['---']*len(b['headers']))+' |'])
            for row in b['rows']:
                result.append('<tr>'+''.join('<td>'+html.escape(str(x))+'</td>' for x in row)+'</tr>')
                md.append('| '+' | '.join(str(x).replace('|','／').replace('\n','<br>') for x in row)+' |')
            result.append('</tbody></table></div><p class="evidence">证据'+html.escape(src)+'</p>');md.append('\n证据'+src+'\n')
        elif k=='image':
            path=ROOT/b['path'];data=base64.b64encode(path.read_bytes()).decode()
            result.append('<figure><img alt="'+html.escape(b['caption'],quote=True)+'" src="data:image/png;base64,'+data+'"><figcaption>'+html.escape(b['caption'])+'</figcaption></figure>')
            md.append('!['+b['caption']+'](../'+b['path']+')\n')
    result.append('</body></html>')
    (OUT/(STEM+'.html')).write_text(''.join(result),encoding='utf-8')
    (OUT/(STEM+'.md')).write_text('\n'.join(md),encoding='utf-8')


if __name__=='__main__':
    blocks=json.loads((OUT/'REPORT_BLOCKS.json').read_text())
    path=docx(blocks);web_and_markdown(blocks)
    print(json.dumps({'docx':str(path),'blocks':len(blocks),'tables':sum(b['kind']=='table' for b in blocks),'images':sum(b['kind']=='image' for b in blocks)},ensure_ascii=False))
