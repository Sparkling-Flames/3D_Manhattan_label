"""Build the Chinese adoption report directly from preserved result tables.

The JSON block source, table extraction ledger, HTML and DOCX contain the same
substantive report. This is presentation code, not an alternative analysis.
"""
from pathlib import Path
import base64
import html
import json
import math
import pandas as pd
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT=Path(__file__).resolve().parents[1]
WORK=OUT.parent
SYN=OUT/'stage05_synthesis';SYN.mkdir(exist_ok=True)
B=WORK/'repo/research/workers/personnel_v21_reselection_20261029'
BLOCKS=[];EXTRACTS=[];SOURCES={}


def source(key,path,desc):
    path=str(path)
    mapping={'audit_classification/':'stage01_classification/','audit_comparison/':'stage02_comparison/','audit_capacity/':'stage02_capacity/','audit_cases/':'stage04_cases/six_cases/','audit_case_vfua/':'stage04_cases/vfua_case/','HOHONET_v21_reaudit_20261029/':''}
    delivery=path
    for old,new in mapping.items():
        if path.startswith(old):delivery=new+path[len(old):];break
    if path.endswith('stage07_adoption/ALL45_METHOD_ADOPTION_DECISIONS.csv'):delivery='stage00_intake/ALL45_METHOD_ADOPTION_DECISIONS.csv'
    SOURCES[key]={'path':path,'delivery_relative_path':delivery,'description':desc}


def read(key):return pd.read_csv(WORK/SOURCES[key]['path'],low_memory=False)


def select(key,**filters):
    d=read(key)
    for k,v in filters.items():d=d[d[k].eq(v)]
    assert len(d)==1,(key,filters,len(d))
    r=d.iloc[0]
    EXTRACTS.append({'evidence_key':key,'filters_json':json.dumps(filters,ensure_ascii=False),'values_json':r.to_json(force_ascii=False)})
    return r


def p(text,src=None):BLOCKS.append({'kind':'p','text':text,'source':src})
def h(text,level=1):BLOCKS.append({'kind':'h','text':text,'level':level})
def table(title,headers,rows,src):BLOCKS.append({'kind':'table','title':title,'headers':headers,'rows':rows,'source':src})
def picture(path,caption):BLOCKS.append({'kind':'image','path':str(path),'caption':caption})
def f(x,n=3,signed=False):
    if pd.isna(x):return '不可评'
    return format(float(x),('+' if signed else '')+f'.{n}f')
def pct(x):return f(float(x)*100,2)+'%'
def ratio(a,b):return str(int(a))+'/'+str(int(b))


source('E01','HOHONET_v21_reaudit_20261029/stage00_intake/INTAKE_AND_COMPUTATION_CONTRACT.json','仓库版本、研究范围与不变合同')
source('E02','audit_classification/results/RAW_GEOMETRY_SUMMARY.csv','普通Manual同点数原答几何，全部参数变体')
source('E03','audit_classification/results/LOBO_STABILITY_SUMMARY.csv','楼外分区相对ALL的描述性ARI')
source('E04','audit_classification/results/DSTRUCT_V21_PARAMETER_AND_RAW_AUDIT_ZH.md','参数追溯、支持门槛、真父子关系及名单')
source('E05','audit_capacity/completed_supplement/FINAL_COMPETITIVE_N_M_FUNNEL.csv','五重点方法四主格及12/4容量—执行—可评漏斗')
source('E06','audit_comparison/abc_integrated/MAIN_CONDITIONS_JOINT_Q.csv','三端同支持、同原A的幅度与方向；背景→块→池→图→楼')
source('E07','audit_comparison/abc_integrated/MAIN_CONDITIONS_METRIC_LONG.csv','Q分项、结构、几何、时间、失败及各自可评分母')
source('E08','audit_comparison/abc_integrated/MAIN_HARD23_COMMON_POOL_READING.csv','TARGET/RAW3共同精确池比较')
source('E09','audit_comparison/conditional_results/CONDITIONAL_METHOD_SUMMARY.csv','相同有限团队的Q/QT条件增量')
source('E10','audit_comparison/conditional_results/LABEL_RANDOMIZATION_REFERENCE.csv','Q层内标签随机置换容量参照')
source('E11','audit_comparison/conditional_results/SAME_COMPOSITION_RESIDUAL_SUMMARY.csv','同配比内部的具体成员残余')
source('E12','HOHONET_v21_reaudit_20261029/stage02_comparison/FINAL_EVIDENCE_REVIEW_ZH.md','旧基线控制、全部来源、父子增量和综合解释；原版与两处显示勘误并存')
source('E13','audit_cases/results/CASE_SUMMARY.csv','六个机制案例的精确Q与几何比较')
source('E14','audit_cases/CASE_MECHANISM_REVIEW_ZH.md','六例原图、原答、109组票账与人工复核点')
source('E15','repo/research/workers/personnel_v21_reselection_20261029/stage07_adoption/ALL45_METHOD_ADOPTION_DECISIONS.csv','原45配置全面比较和本轮保留的历史基线')
source('E16','HOHONET_v21_reaudit_20261029/stage05_synthesis/PERSONNEL_ROSTER_ALL_WITH_HISTORY.csv','62名人员ALL描述名单与历史支持；另有全楼外标签')
source('E17','audit_capacity/completed_supplement/COMPLETED_SUPPLEMENT_AUDIT.json','6548新事件、16818端点与独立核验')
source('E18','audit_comparison/abc_integrated/MAIN_CONDITIONS_DIRECTION_PATTERNS.csv','0.1/1/2Q近零带的同向、反向与单边响应')
source('E19','audit_capacity/completed_supplement/FINAL_REPLACEMENT_BAND_COVERAGE.csv','固定m、半队、过半与较大队的最终独立覆盖口径')
source('E20','audit_case_vfua/VFUA_COMMON_POOL_NEGATIVE_TAIL_ZH.md','VFua共同池负尾、两套ABC六端点、110组票账与同Z背景响应')
source('E21','audit_capacity/results/ALL_N_M_DIRECTION_Q_RESULTS.csv.gz','既有大团队直接AC的原事件→池→图口径，不混入主ABC块等权')
source('E22','audit_comparison/abc_integrated/ALL_STAGES_DIRECTION_N_M_SUMMARY.csv.gz','全部旧新阶段的精确方向与人数汇总；TARGET12/4和RFUSE26/3的直接来源')
source('E23','HOHONET_v21_reaudit_20261029/stage05_synthesis/ADOPTED_CLASS_MEMBERS_AND_PROFILES_ALL.csv','采用方法的ALL完整成员和历史残差画像；分类输入与后验Q分开解释')


def build_content():
    h('HOHONET v2.1 人员分类独立复核',0)
    p('普通 Manual 为主｜分类形成、有限配比、严格 ABC、大块替换与真实机制案例｜2026-10-29')
    p('读取仓库提交：519c7447278dd16c3d69f505b807fb5b53792b36。既有 v2.1 研究冻结来源与本轮追加结果分别保留；本报告没有覆盖旧采用表、原答、参考或范围裁定。','E01')
    h('一、可实际采用的判断')
    p('推荐普通 Manual 以 TARGET_DSTRUCT_K3 为人员分类主方法，以同一层次树的 K2 为粗层，K4仅用于确有父内差异与足够容量的局部研究。OPPJOINT K3/K4承担联合画像与有方向的配比组织；RFUSE_GEOM_K2保留融合响应反推互补。本轮重调并验证的 V21_DSTRUCT_RAW3_K3，用作不含参考对应轴的定义敏感性和八人替换扩覆盖方案。这个优先级是重新审查全部既有证据和本轮补算后的研究判断，不沿用“OPPJOINT已经是主分类”的安排。','E04/E06/E08/E12')
    p('“人员画像用OPPJOINT、人员分类回到DSTRUCT”的分工有充分依据，但具体落地需要两个限定：主分类采用当前v2.1目标对齐的TARGET版本；调参后容量更宽的RAW3有明确增量，却没有在原答解释、稳定性、七人覆盖及共同池表现上形成一致优势，因而不直接替代主方法。')
    table('采用顺序与用途',['方法/层级','正式用途','决定性依据及边界'],[
      ['TARGET_DSTRUCT K3；K2粗层','解释不同人群的原答结构响应；主混合配比与严格ABC','标法语义、同点数原答差异、楼外稳定性、七人覆盖及正负方向共同支持；含参考匹配诊断，不能称完全无参考。'],
      ['TARGET_DSTRUCT K4','局部父内细分','额外四人小类可研究；没有整体几何分离优势，不自动升为主层。'],
      ['OPPJOINT K3/K4','联合人员画像、质量/结构/位置/时间配比','K3承接主要质量角色；K4突出慢小类和时间构成。主升质方向不属于第四类的独有增量。'],
      ['RAW3 K3','纯原答定义与八人容量补充','剔除unmatched_ref_9后扩大八人合法池；类别语义变宽、七人覆盖与稳定性下降。'],
      ['RFUSE_GEOM K2','由融合响应解释人员影响','保留正反向互补；少数类四人，适合三至四人块，不承担七八人纯类主实验。'],
      ['OPPSTRUCT / OPPGEOM / MSDGEOM_SCOPE','条件定义、几何机制、跨Manual/Semi及scope补充','已有比较和具体用途保留；不将输入域不同或未获主位等同于无效。']],'E04/E12/E15')
    p('后续正文以TARGET K3简称“T”，RAW3 K3简称“R”。简称只为表格排版；所有结果仍使用完整method、heldout_building、pool_id及事件键。类别是依历史和问题定义形成的群体倾向，不是永久能力等级。')

    h('二、当前数据域、既有材料与本轮新增')
    table('当前合法池入口',['数据域','池数 / 独立图片 / 建筑','本轮地位'],[
      ['普通Manual主目标','206 / 206 / 22','主研究与大块严格ABC'],
      ['普通Manual其他准入目标','30 / 28 / 9','同一图片中的其他合法目标；人数3–17，不能支持本轮七八人纯类严格ABC'],
      ['全部普通Manual','236 / 206 / 22','236是目标池数，不是236张互不重复全景图'],
      ['门洞Manual；特殊Manual','17/17/9；2/2/2','保持独立域，不并入普通Manual结论'],
      ['普通Semi；特殊Semi','38/37/12；1/1/1','补充；新增其他准入Semi仅4人']],'E01/E05')
    p('已经完成的材料包括完整有限配比、45配置广筛、重点ABC、固定原队扩大替换块、固定替换块改变背景或人数、大团队AC，以及原答和票账案例。98是数据域×配置条目，不是98种独立分类思想；45也包含同家族粗细、轴基线、目标权重与反向变体。RFUSE_GEOM_K2与TARGET_RFUSE_GEOM_K2在ALL和22楼外折同分区，不能算两套独立成功思想。','E12/E15')
    p('本轮实际新增五种特征/参数变体×K2/K3/K4，共15个配置，另精确复现原TARGET三个配置；每套都做ALL与22个留目标建筑外折。新增严格ABC共6,548个方法事件，16,818个去重团队端点执行；其中8个是旧标量已有但几何未在当前缓存中的必要重放，真正新增物理团队16,810。没有新人员、新房间、新同房研究设计，也没有重跑256条自然人数曲线。','E04/E17')
    p('四个主要格为n14/m7（七人半队）、n20/m7（较大队七人）、n12/m8（八人过半）、n16/m8（八人半队）。既有方法的m7只补过去没有登记的方向×池格；m8用统一两块成员与可行low/high背景建立真实嵌套路径。RAW3另加n12/m4作宽覆盖桥接。338个已有m7格没有为证明存在而重算。')
    p('各类画像和变体配置在读取本轮新团队结果前冻结；最终采用则明确使用完成后的全证据选型。楼外形成标签避免把目标楼直接用于人员特征拟合，但同一研究数据上的方法选择仍不是全新外部验证，不能报告为预测准确率。[方法文献M3]')

    h('三、先前怎样构造、比较和选择；这次修正什么')
    p('既有路线先从个人历史形成结构、边界/尺度、质量与时间等域，按记录/目标/图片/建筑处理残差和等权，再在每个楼外折拟合稳健缩放、域内降维与层次聚类。原答前向路线从实际标注特征出发；机会路线用当前参考可计算的结构条件与响应；联合路线同时纳入Q/T等画像；RFUSE反向路线从自然融合响应描述人员影响。随后用原答对距离、有限构成、广筛与重点替换作比较。','E04/E12/E15')
    p('先前选择OPPJOINT为主，主要依赖其联合解释、醒目的换入升Q方向、时间小类和较完整重点展示。这些实测结果值得保留，但对于本轮“不同人群怎样标、怎样改变共识”的目标，不能将最大升Q方向放在标法语义与多条件覆盖之前。重点扩展数量不同，也会使某路线看起来证据更丰富；有容量而未进入设计应先补齐，不能当作方法没有作用。')
    p('本轮采用顺序为：先检查历史域和标法可解释性；再检查同点数原答差异与楼外稳定性；再看真实池容量及同A的同类/跨类变化；最后以Q分项、几何、记录劳动量、失败和真实票账解释。没有建立一个把这些指标任意加权成单分数的排名。主方法必须在核心问题上有清楚优势；互补方法要有明确新增问题和适用条件。')
    p('最重要的对照修正是旧Q基线：OPPJOINT K4与Q二分虽有56共同池，精确相同D为0、相同A也为0。55池共同历史限制并未匹配连续历史Q强度；联合换入Z相对X的楼外历史Q差约16.962，Q二分为10.613。原来的约+8.05与+3.93可以说明不同选队规则的实际结果，不能单独证明联合结构/时间的额外作用。K3真父方向与K4有142个事件共享D/A/X/Y，说明主要升质角色在父层已存在。','E12')

    h('四、DSTRUCT的v2.1参数是否合适')
    table('沿用的规则与已重算的量',['环节','延续的设计','v2.1已做的更新'],[
      ['特征','turn_n、short_turn_edges、near_straight_nodes_5deg、unmatched_ref_9','当前原答与当前目标参考诊断；unmatched_ref_9同时受缺漏和定位影响'],
      ['条件域','总体；困难与困难−简单','正式难度已由v2.1 current_class更新；TARGET不把其他空间等级复制给未知新目标'],
      ['历史范围与权重','逐层等权、相对任务同行的残差','TARGET去除111条特殊范围辅助历史；普通合法目标对齐，记录→目标→图片→建筑'],
      ['支持门槛','总体12图3楼；简单5图2楼；困难4图2楼','每折重新计算支持；普通Manual实际原答6,739条'],
      ['缩放/聚类','median/(p85−p15)，退化std或1；clip±3；域内PCA最多2维，域RMS等权；Ward单树切K2/3/4','每个楼外折重算中位数、尺度、PCA、树与成员；不是沿用旧名单或旧数值']],'E04')
    p('因此“沿用旧难度分层思想”和“还在使用旧版难度数值”是两件事。当前没有查到后者；本轮质疑并实际检验的是条件域、参考诊断轴、维数截断、域权重和支持阈值是否适合本研究。TARGET不直接输入Q或时间，但不能因此宣称完全不依赖参考。')
    table('实际完成的调参矩阵',['变体','具体改变','本轮判断'],[
      ['FULL','域内保留全部标准化维度，取消PCA2截断','原答均值略增但建筑等权与稳定性未改善，不足以换主方法'],
      ['RAW3','删除各层unmatched_ref_9，仅保留三条原答表达轴','定义更纯；八人容量扩大，作为正式补充'],
      ['LEVEL','只保留总体结构域','检验困难条件的额外信息；未优于TARGET'],
      ['COND','以可计算challenge及challenge−plain替代人工困难域','条件定义改变后没有一致优势'],
      ['HALF','困难域相对总体域从1:1调为1:0.5','检验困难响应是否权重过大；未产生替代证据']],'E04')
    p('机会条件基于参考非冗余转折、边界隐藏顶点数H与短边；H不是房间高度。OPPSTRUCT的连续机会斜率已经存在，不是本轮新提出。COND是针对人工难度条件的另一种实际检验。支持门槛从总体12图3楼放宽到8图3楼，23折都仍只有59名可分类者；W019/W026普通Manual历史为0，W022仅2图1楼。缺少七八人容量主要来自真实池中的类别分布，不能归咎于这个门槛。','E04')
    rows=[]
    for method in ['TARGET_DSTRUCT_K3','V21_DSTRUCT_RAW3_K3','V21_DSTRUCT_FULL_K3','OPPSTRUCT_K3','OPPJOINT_K3','OPPJOINT_K4']:
        r=select('E02',method=method,support='common_old45_new',subset='same_corner_count',scope_summary='ALL_TARGETS',difficulty_summary='ALL',metric='geometry_pair_angle')
        stability=select('E03',method=method)
        rows.append([method.replace('_',' '),f(r.gap),ratio(r.positive_gap_images,r.images),ratio(r.positive_gap_buildings,r.buildings),f(stability['mean'])])
    table('同点数原答的角距差与楼外分区一致性',['配置','类间−类内角距°','正差图片','正差建筑','LOBO/ALL ARI'],rows,'E02/E03/E04/E15')
    p('同点数限制避免仅由点数多少制造差异。本表覆盖普通Manual全部合法目标，同图多目标等权。就TARGET与RAW3 K3而言，其他准入目标贡献7池、6图；各方法可评目标略有差异，完整分母见逐池表，不能误称206图全部仅有主目标的同等原答对支持。TARGET与RAW3在相同206图上配对后，RAW3相对TARGET的角距差平均仍减少约0.248°；FULL仅增加约0.019°，且建筑等权差略负。ARI是分区一致性指标，不是人员类型真值准确率；0.885不能读成88.5%的分类正确。[方法文献M1、M2] 本轮有针对性的15配置检验支持保留TARGET当前参数，不证明它是所有可能算法中的唯一最优。','E04')
    p('RAW3在B6ByNegPMKs留楼外折中，低结构两类在hard与challenge轴上的次序交换。本报告TARGET/RAW3的主方向按困难域转折残差秩，OPPSTRUCT按challenge秩，OPPJOINT按历史Q秩；两种结构秩的转换表完整保留。更改角色命名不删除事件，也不改变所有方向合并的容量。','E12')

    h('五、同配比下换具体成员，剩下多少变化')
    p('原有限面板固定D，从12名变量候选中选6人。共同历史资格筛选后，182个面板中有174个保留12名、1个保留11名、7个保留10名可用变量人员，完整保留相应六人组合，共162,708个面板—团队关联，覆盖38图22楼；所有方法使用相同有效候选与团队，Q随后因评分缺失降至161,580。这是有限候选池的完整组合，不是全体人员任意团队的总体穷举。每一构成格的团队数、均值、SD、10/25/50/75/90分位和极值均保存；仅一个有效团队的格不伪造SD或IQR。','E12')
    rows=[]
    for method in ['TARGET_DSTRUCT_K3','TARGET_DSTRUCT_K4','OPPSTRUCT_K3','OPPJOINT_K3','OPPJOINT_K4','V21_DSTRUCT_RAW3_K3','RFUSE_GEOM_K2']:
        r=select('E11',method=method,metric='B8L2',weighting='building_equal')
        rows.append([method.replace('TARGET_DSTRUCT','TARGET').replace('V21_DSTRUCT_RAW3','RAW3').replace('_',' '),pct(r.within_share),f(r.within_rms_sd)])
    table('同配比内部的Q变异仍很大',['方法','仍在同配比内部的Q变异','格内Q均方根SD'],rows,'E11/E12')
    p('这些SD是按既有层级汇总格内方差后开方，不能说每一张图都变化6分。结论是类别构成组织了一部分变化，具体成员仍然是大的变异来源。主实验必须保留同类B；同类并不意味着可交换的复制品。细分类使格数增多，自然可以提高样本内解释比例，因此必须结合真父子、容量与实际替换。')

    h('六、固定七人、固定八人、半队和大队：完整漏斗')
    p('严格纯类ABC为A=D∪X、B=D∪Y、C=D∪Z，四块真实互斥，X/Y同类且各m人，Z为另一类m人，每队n人。容量条件为N≥n+2m、源类人数≥2m、目标类人数≥m。未知历史人员可以留在D，不能进入已知纯类X/Y/Z。直接AC没有B，另要求N≥n+m、源与目标类各≥m。前者的不足不等于所有混合构成块都不可能。')
    p('全部普通Manual入口236池；本轮四主格中，30个其他准入池均在总人数层退出。n14/m7与n12/m8各178池总人数足够；n16/m8为150池；n20/m7为120池。五套重点分类在这些池的已知人数总量均达到3m，继续退出来自类别分布。本表每格为“类别容量允许＝已执行 / C−B可评 / 三端幅度可评”；不可评数不是零分。','E05')
    cap=read('E05');rows=[]
    for m in ['TARGET_DSTRUCT_K3','V21_DSTRUCT_RAW3_K3','OPPSTRUCT_K3','OPPJOINT_K3','OPPJOINT_K4']:
        vals=[]
        for n,k in [(14,7),(12,8),(16,8),(20,7)]:
            r=select('E05',method=m,n=n,m=k)
            vals.append(f"{int(r.class_feasible_pools)} / {int(r.Q_CB_evaluable_pools)} / {int(r.Q_gap_evaluable_pools)}")
        rows.append([m.replace('V21_DSTRUCT_RAW3','RAW3').replace('TARGET_DSTRUCT','TARGET').replace('_',' ')]+vals)
    table('容量到结果的最终覆盖（池数）',['方法','14/7 半队','12/8 过半','16/8 半队','20/7 较大队'],rows,'E05')
    p('这四个网格、五套重点分类的“容量有但仍未登记”均为0。并不表示所有可能n/m全部执行，也不表示每格枚举了所有成员组合。RAW3的八人半队容量64池19楼，高于TARGET的41池15楼；七人半队则118池22楼，低于TARGET的128池22楼。不能只报ALL人数，也不能只报最有利方向的容量。')
    p('各池最大35人，n20/m8严格ABC至少需要36人，因而无论分类怎样调都不可行。较大团队AC可以合法，但没有B，不能据其升质推断跨类幅度超过同类。已有OPPJOINT K4历史Q秩1/4→4/4直接AC：n28/m7为30图，C−A约+2.68、29图为正；n30/m5为64可评图/65登记图、约+2.50、55图为正；n32/m3为76/77图、约+1.48、64图为正。这些沿用原事件→池→图的AC汇总，单列为大队直接替换证据，不与主ABC的XYZ先背景平均口径混称。','E21')
    rows=[]
    for band,name in [('fixed_m7_any_n','固定m7，任意n'),('fixed_m8_any_n','固定m8，任意n'),('exact_half_all_m','恰半队，全部m'),('exact_half_m_ge7','恰半队，m≥7'),('above_half_all_m','过半，全部m'),('above_half_m_ge7','过半，m≥7'),('large_team_ABC_n_ge20','n≥20严格ABC'),('large_team_AC_n_ge20','n≥20直接AC')]:
        vals=[]
        for method in ['TARGET_DSTRUCT_K3','V21_DSTRUCT_RAW3_K3','OPPSTRUCT_K3','OPPJOINT_K3','OPPJOINT_K4']:
            r=select('E19',method=method,population='all_ordinary_236',band=band)
            vals.append(f'{int(r.class_feasible_pools)}/{int(r.executed_pools)}/{int(r.primary_Q_evaluable_pools)}')
        rows.append([name]+vals)
    table('全部236普通Manual池：不同问题的容量/执行/成对Q可评并集',['问题','TARGET3','RAW3','OPPSTRUCT3','JOINT3','JOINT4'],rows,'E19')
    p('数学枚举限定n≥3、1≤m≤n、当前池N≤35；过半包含m=n的全队替换，逐池非空D容量另存。固定m7的全局合法n为7–21，固定m8为8–19；区间内每格仍需逐类容量检查。数学范围允许D为空，不能冒充已选定的展示设计；本轮四主格均保留非空D。ABC成对Q指C−B，AC指C−A；并集只要求该池某一规模/方向有可评事件，不表示全部格或成员组合都可评。n≥20行包含小块替换；若再要求m≥7，严格ABC只可能m7、n20或21，其可行池并集已由本轮n20/m7覆盖。','E19')
    p('任意n的固定m7仍有三个方法×池未登记：RAW3的uNb9QFRL6hY-66（N27，可n7–13），OPPJOINT K3/K4的q9vSo1VnCiC-29（N26，可n7–12）。它们在n14/m7本就总人数不足，不影响本轮四主格完整性；也没有证据表明只补这两张较小队图片会改变主方法或七人半队适用判断，所以没有把它们包装为本轮必要全量补算。全部其他未设计n/m格仍明确列在逐池表。','E19')
    p('对30个其他准入小池，TARGET小m恰半队有12池容量、6池执行且可评；过半有14池容量、7池执行、6池可评。其余方法的同口径行完整保存。它们不是因为七八人不足就全部退出人员模拟研究。','E19')

    h('七、实际换人作用：同一原队A、同时报告幅度与方向')
    p('主幅度W=|Q_B−Q_A|，C_A=|Q_C−Q_A|，G=C_A−W；有符号方向分别报告C−A与C−B。表内这些Q量都要求A/B/C共同可评分；更宽的A/C或B/C成对分母保留在原表，不混接。旧对称跨幅度(|C−A|+|C−B|)/2也保留供回连，但不再与同A口径混称。')
    p('聚合先在固定XYZ内对背景平均，再池内对块平均、同图对目标平均，最后图等权；建筑均值另存。m7合并旧与本轮漏格补算；m8使用本轮统一设计。所选网格有12,932旧方法事件与6,548新方法事件，共19,480条；相关事件不是独立图片样本。','E12')
    rows=[]
    for n,m in [(14,7),(20,7),(12,8),(16,8)]:
        stage='new_only' if m==8 else 'combined_union'
        for method,name in [('TARGET_DSTRUCT_K3','T'),('V21_DSTRUCT_RAW3_K3','R')]:
            r=select('E06',method=method,n=n,m=m,axis='hard_turn_rank_asc',rank_direction='2/3->3/3',stage_scope=stage,scope='primary')
            rows.append([f'{name} {n}/{m}',f"{int(r.images)}/{int(r.buildings)}；Q{int(r.joint_Q_CA_valid_images)}",f(r.B8L2_same_origin_within_mean_image,2),f(r.B8L2_same_origin_cross_mean_image,2),f(r.B8L2_same_origin_gap_mean_image,2,True)+'；'+ratio(r.B8L2_same_origin_gap_positive_images,r.B8L2_same_origin_gap_valid_images),f(r.joint_Q_CA_mean_image,2,True)+'；'+ratio(r.joint_Q_CA_positive_images,r.joint_Q_CA_valid_images),f(r.joint_Q_CB_mean_image,2,True)+'；'+ratio(r.joint_Q_CB_positive_images,r.joint_Q_CB_valid_images)])
    table('由中等表达角色换向较多转折角色（hard秩2→3）',['方法 n/m','登记图/楼；Q图','同类W','跨类C_A','G；正图','C−A；正图','C−B；正图'],rows,'E06')
    p('主结构展示优先TARGET n14/m7；较大队七人替换用TARGET n20/m7；八人展示把RAW3作为扩覆盖补充。TARGET上行n14/m7的G建筑等权为+2.578，n20/m7为+1.236；RAW3对应+1.409与−0.150。RAW3 n20/m7虽多数图C−A为正，跨类绝对幅度相对同类的图均值仍为负，不能泛称所有大块格都有额外区分。','E06/E12')
    p('已有宽覆盖TARGET n12/m4上行方向有179登记图22楼、178三端Q可评图；按本轮去重及背景/块等权，W约5.181、C_A约6.321、G约1.140，127/178图G>0，C−A约+4.675。它适合先展示广覆盖再进入七八人条件。该估计包含已有多阶段块，与旧仅重点阶段的均值不必相同。直接来源过滤为combined_union/primary/TARGET_DSTRUCT_K3/n12/m4/hard_turn_rank_asc/2/3→3/3。','E22')
    rows=[]
    for n,m in [(14,7),(16,8),(20,7)]:
        for method,name in [('TARGET_DSTRUCT_K3','T'),('V21_DSTRUCT_RAW3_K3','R')]:
            r=select('E06',method=method,n=n,m=m,axis='hard_turn_rank_asc',rank_direction='2/3->1/3',stage_scope='new_only' if m==8 else 'combined_union',scope='primary')
            rows.append([f'{name} {n}/{m}',str(int(r.joint_Q_CA_valid_images)),f(r.joint_Q_CA_mean_image,3,True),ratio(r.joint_Q_CA_negative_images,r.joint_Q_CA_valid_images),f(r.joint_Q_CB_mean_image,3,True),f(r.B8L2_same_origin_gap_mean_image,3,True),ratio(r.B8L2_same_origin_gap_positive_images,r.B8L2_same_origin_gap_valid_images)])
    table('向较少转折角色的负方向也保留（hard秩2→1）',['方法 n/m','共同Q图','C−A','C−A负图','C−B','G','G正图'],rows,'E06')
    p('TARGET14/7向较少表达角色，旧少量登记曾出现均值为正而多数图为负的结果；补齐容量缺口后，46张三端可评图的C−B约−1.811。这个修正说明既有重点方向和少量支持不能代表全可行范围，也不需要把负方向从分类价值中删除。TARGET16/8的负方向确实降Q，而G略负；RAW3该格降Q更明显且G为正。二者回答不同问题。','E06/E12')
    p('响应形态不只是“同类作用放大”。以±1Q为描述性近零带，TARGET14/7上行83图中，同向28、反向20、仅跨类28、仅同类3、双近零4；20/7的70图为23、13、20、5、9。完整±0.1与±2Q敏感性同时保留；阈值未经验证为最小重要差异。','E18/E12')

    h('八、共同图片上的方法比较与背景依赖')
    rows=[]
    for n,m in [(14,7),(20,7),(12,8),(16,8)]:
        r=select('E08',stage='new_only' if m==8 else 'combined_union',n=n,m=m,role='hard_turn_rank_asc:2/3->3/3',metric='B8L2_same_origin_gap')
        rows.append([f'{n}/{m}',str(int(r.common_registered_pools)),f'{int(r.images)}/{int(r.buildings)}',f(r['T'],3,True),f(r['R'],3,True),ratio(r.Tpos,r.images),ratio(r.Rpos,r.images)])
    table('相同目标池、n/m与结构角色后的Q额外幅度',['n/m','共同登记池','共同Q图/楼','TARGET G','RAW3 G','TARGET正图','RAW3正图'],rows,'E08')
    p('这控制了目标池和规模，但通常仍未控制两法完全相同的A、D、具体成员或连续历史Q，不能称分类的因果独立效应。m8共同七池中六池的全部实体ABC相同，输出相同不是独立复现；只有pa4otMbVnkk-05的实际选队不同。RAW3八人上行从TARGET的11图扩至29图，主要新增的是可组队范围，不能包装成对同一批图普遍更强。','E08/E12')
    p('n16/m8七共同池里六图G为正，均值却为负，主要负尾为VFuaQ6m2Qom-04：一组三队Q为79.7922/26.1511/44.8821，C−B=+18.731但C−A=−34.910；另一组背景三队均不可评。两法实体团队相同。此图需用原答和票账解释，不能简单说“RAW3参数错误”或“有反例正常”。第七案例单独保留完整回查。','E12')
    p('相同XYZ、D嵌套的已有路径也显示人数与背景改变会改变净响应。原联合m6从n12到n20的17共同图，C相对B的净变化平均约−2.567，仅5图增加；TARGET m7从n10到n18的18共同图约−1.07，仅5图增加。添加背景人员也提高多数门槛，现有观察不能把“背景内容”与“人数/门槛”单独因果拆开。本轮m8同XYZ n12→16的逐事件四端共同可评表已保存，继续同时报告C、B和(C−B)的变化。','E12')

    h('九、Q分项、几何、记录劳动量、失败分别解释')
    p('TARGET14/7上行的成对C−A，D10约+6.222、IoU3D×100约+3.958、底界角误差−0.445°、转折+0.378、短边+0.148；每项采用自身可评分母。配对角距的同A额外变化约+0.208°，53/83图为正；BEV距离额外变化约+0.0136，45/83图为正。几何的多数区分弱于Q，不能把升Q扩写成每个几何指标都更可分。','E07/E12')
    rows=[]
    for method,name in [('TARGET_DSTRUCT_K3','T'),('V21_DSTRUCT_RAW3_K3','R')]:
        for direction in ['2/3->3/3','2/3->1/3']:
            for metric,label in [('D10_CA','D10 C−A'),('IoU3D100_CA','IoU3D×100 C−A'),('fusion_turn_n_CA','转折数 C−A'),('geometry_pair_angle_same_origin_gap','角距G（度）'),('geometry_bev_distance_same_origin_gap','BEV距离G'),('record_time_CA_pct_source_matched','记录劳动量C/A变化%')]:
                r=select('E07',method=method,n=16,m=8,axis='hard_turn_rank_asc',rank_direction=direction,stage_scope='new_only',scope='primary',metric=metric)
                rows.append([name+' '+direction.replace('/3',''),label,f(r.mean_image,3,True),ratio(r.positive_images,r.valid_images),f(r.mean_building,3,True)])
    table('八人半队的多指标结果（各项分母独立）',['方法/方向','指标','图片均值','正图/可评图','建筑均值'],rows,'E07')
    p('正图表示该数值大于0，未自动等同“更好”：误差和损失通常越低越好，转折更多未必更符合当前范围，几何距离变大也不代表质量改善。本轮16/8的TARGET上行角距G平均为负但6/11图为正；RAW3接近零但20/29图为正，且建筑均值更负，说明少数大响应与楼别分布必须保留。')
    p('来源人数向量匹配的记录劳动量，TARGET14/7 C相对A平均+7.994%（56图），20/7为+4.751%（42图）；RAW3为+10.515%（33图）与+5.674%（25图）。这不是新的现场计时，也不是未来生产耗时预测。在上述hard2→3上行方向，m8的C/A来源匹配时间仅TARGET三图、RAW3两图可评，不能扩写为全图片规律；不同日志来源、来源人数构成和缺失均单列。','E07/E12')
    p('全部16,818个本轮执行端点中，16,626可评Q，192不可评：129个no_connection、62个below_four、1个invalid_floor。缺失Q不填0。新16/8上行TARGET22个事件的A/B/C不可评计数为1/1/2，RAW3 76事件为2/2/3；这些是事件端点计数，图片等权失败率另表。结构或几何仍可计算的失败端点只保留该项合法结果，不伪造完整环。','E17/E07')
    p('旧TARGET20/7固定A部分只保存可回连的标量，当前交付没有对应成对角距/BEV多边形值；该部分保持缺失。本轮m8保存全部端点几何。来源匹配时间和几何缺失不应通过换分母悄悄变成方法优劣。','E12')

    h('十、联合分类相对Q/QT到底增加什么')
    p('本轮不重复旧Q二分选队实验，而在相同182面板、同一D、同一候选人、同一精确团队和已保存端点上做条件比较。第一层固定Q高/低完整人数向量；第二层加入连续历史Q均值、目标等权Q、低尾Q、二阶项；第三层加入历史log时间及Q×时间项。候选完整离散构成格与Q人数格交叉加入模型。同一指标内所有方法共享支持，跨指标不同：Q为182面板/38图/22楼、161,580关联；转折可评162,315关联；时间为130面板/31图/21楼、116,966关联。所有分析只是有限样本内描述性剩余方差，不是预测、因果独立效应或升质百分比。','E09/E12')
    rows=[]
    for method in ['TARGET_DSTRUCT_K3','OPPSTRUCT_K3','OPPJOINT_K3','OPPJOINT_K4','RFUSE_GEOM_K2','V21_DSTRUCT_RAW3_K3','V3_JOINT_QT_COARSE']:
        vals=[]
        for metric,model in [('B8L2','exact_Q_quota'),('B8L2','quota_continuous_Q'),('fusion_turn_n','quota_continuous_Q'),('record_time_single_source','quota_continuous_QT')]:
            r=select('E09',method=method,metric=metric,model=model,weighting='building_equal');vals.append(pct(r.candidate_extra_share_total))
        rows.append([method.replace('TARGET_DSTRUCT','TARGET').replace('V21_DSTRUCT_RAW3','RAW3').replace('_',' ')]+vals)
    table('固定候选、背景、端点后的额外总变异解释（建筑等权）',['分类','Q配比控制后：Q','连续Q控制后：Q','连续Q控制后：转折','Q/T控制后：时间'],rows,'E09')
    p('Q基线本身解释11.36%（仅二分配比）、33.73%（连续Q）及42.09%（再加时间）的Q变异；Q/T模型已解释单一来源记录时间的89.14%。因此OPPJOINT K4无条件时间构成解释43.63%，不能全部算成超出Q/T的新增作用。不同条件模型的增量不必单调，应在同一基线上比较合计解释和剩余量。','E09/E12')
    p('为了诊断更多构成格自然提高拟合的问题，另外在各面板Q层内随机置换候选类别128次，保持每层类别容量和D不变。以下随机参照只固定Q二分人数，未加入连续Q/T：TARGET K3的Q增量18.12%对随机均值11.90%，转折20.14%对11.87%；OPPJOINT K4时间增量39.16%对23.72%，TARGET时间13.63%对13.67%。这支持结构分类用于结构，联合画像用于时间；也说明不能把全部样本内增量当成方法特有作用。随机参照不是总体p值，面板之间不是独立样本。','E10/E12')
    p('不需要证明结构作用完全独立于Q才承认研究价值。本轮结果支持“在已控制的历史Q/QT表征之外仍有可描述的构成信息”，同时明确剩余Q信息、格数自由度和未控制的非线性仍可能存在。旧V3联合在整体配比与时间组织上仍有较高解释，但其五类不能凭一个最高比例取代更清楚的结构主方法。')

    h('十一、粗细层、正反向与跨条件的实际分工')
    p('TARGET K2为47/12人；K3把47人父类拆成10人简化/近共线较多组和37人常规组；K4再从常规父类拆出W006、W008、W029、W034四人小类。K3到K4的同点数原答角距差约从0.684°降为0.536°，正差图从162/206变154/204，故细分没有自动获得更强主解释。全部父子映射与每个楼外折的成员已保存。','E04/E16')
    p('OPPJOINT K3为19/29/8人，K4将29人父类拆出慢三人W027、W036、W121；主八人高Q角色保留。已有13,544个严格父内ABC，不应再提为未完成。慢小类的8,316个无序m1 ABC只覆盖9图8楼，主要由W121支撑，不能说三人都在同样图片上产生同样效应。K4八人高Q方向本轮仅1图1楼，K3相应方向11图5楼；联合八人展示优先K3，K4细画像保留。','E12/E16')
    p('RFUSE_GEOM_K2由自然融合响应反推人员，ALL为55/4人，小类W109、W112、W126、W130。它的同点数原答角距均值差仅约0.017°，不宜写成最明显的可见原答大类；但已有n26/m3方向覆盖112登记图、111三端Q可评图，按本轮合并口径同A额外幅度约+0.724，68/111图为正，共同C−A约+1.501、C−B约+1.916。它回答谁在融合中怎样发挥作用，保留了与原答前向分类的互补性。直接来源过滤为combined_union/primary/RFUSE_GEOM_K2/n26/m3/response_hard_turn_rank_asc/1/2→2/2。','E22/E15/E16')
    p('已有RFUSE加入Q、目标等权和K3变体均保留。目标等权GEOM K2没有改变任何折的分区；其他反推变体有容量或响应变化，但现有证据不需要为了名称对称另造反向分类。反推K3较大的七八人容量不自动使其成为普通Manual原答主方法。')
    p('Manual/Semi与scope路线保留为域补充。MSDGEOM_SCOPE_K4的ALL人数4/42/4/9，已有尺度和重点ABC用途；其较高配比解释不能改写为普通Manual内的同一问题。Semi自身已有联合粗/细配比Q解释约21.48%/23.16%，结构约6.35%/11.85%；这些来自Semi自身标签和共同支持，不能从Manual45结果直接外推。其他准入目标的有限配比也有作用，但仅30池28图9楼，需与主目标结果分别陈述。','E12/E15')

    h('十二、反例的具体机制与可以直接看的图')
    table('六个冻结机制案例；均为本轮n16/m8严格ABC',['图片 / 事件前缀','已核实机制','结论及人工问题'],[
      ['7y3sRwLe3Va-04 / 2978172','四角均16票，Q约95.5；AC角距0.347°大于AB0.226°','高分相似结构下Q微变而位置仍变；本例无需新范围裁定'],
      ['B6ByNegPMKs-48 / 8ffd905','三臂四角满票，Q 51.33/63.03/54.65；共同顶边下移','不能解释为高分集中；需看通道两端、主顶面/梁底与远端定位'],
      ['uNb9QFRL6hY-40 / 7881ea8','同类B使缝边远角从7票到10票；C仅6票','Q 48.44/97.31/48.43；类内成员与阈值共同放大，需分清近/远两边界'],
      ['x8F5xyUWy9e-08 / ce5e6ff','两种局部角各8票、人员互斥，无来源环覆盖全部入选组','C不可连接；人工看x≈144/155处是否同角定位差或不同细部'],
      ['jtcxE69GiFV-13 / 0836672','短折达到8/9票；共同D的W106观测发生对应重分组','Q约74.43/75.27/92.87；升质伴来源匹配记录劳动量+12.58%'],
      ['S9hNv5qa7GM-09 / bd186f9','当前停止边和两个旧分支角各8票，另一分支7票未过','当前四角范围已裁定；C连接失败，不能把旧分支当正确细化，也不重复裁定']],'E13/E14')
    p('六例按完整结果后的机制选择，不代表总体胜率。完整18臂full_state与原端点标量最大差1.42×10⁻¹⁴；109组、1,285条观测票账保留人员、原答坐标和来源组。每人每组最多一票；n16门槛为8，半票也会入选。局部node编号不能跨臂直接对齐，追踪使用object_id与source_pair_index。固定D只固定原答，不保证加入Z后对应分组不变。','E14')
    picture(WORK/'audit_cases/figures/0836672f35ad7742021ab716/ERP_BEV.png','图1．jtc13：同类B保持四角，跨类C形成六角。虚线为当前目标参考；细线是真实原答，粗线为真实融合。几何按相机高归一，不是未知的米制尺度。完整票账显示补票与对应调整同时发生。')
    picture(WORK/'audit_cases/figures/7881ea8d8b010028bc505a47/LOCAL_CRITICAL_VOTES.png','图2．uNb40水平缝附近。D对远角固定四票，X/Y/Z分别贡献3/6/2票；只有同类B超过门槛8。近方位不等于同一远近结构，真实结构与遮挡下落点仍需人工核查。')
    p('第七例VFuaQ6m2Qom-04专门解释共同七图中均值与多数比例相反。首组A/B/C为4/5/9角：A的粗四角保留大体周边；同类B使右侧粗角由8票降至6票，邻近细角仍只有7/7/7票，长对角线截去大幅底面，遗漏比例由0.1287升至0.5183、方向误差由约0.58°升至21.92°。C恢复一部分细角，却仍遗漏主折与多处7票细角，遗漏比例0.2577、方向误差13.58°。因此C−B为正，C仍比A低34.91分。来源匹配A/C记录劳动量由1872.385秒增至2257.643秒（+20.58%），同时Q下降34.91分。','E20')
    p('VF另一组三队均不可评，仍完整记账：A/B入选组没有共同支持的完整来源环，B入选角最低支持9票，说明失败不都来自恰半票竞争。两组C的换入Z八人相同；背景改变使粗右角x≈635的D票由3增至5、Z保持3票，从6票未选到8票入选，细角x≈691也有另一组互斥8票，从而触发连接失败。这是固定Z对背景响应的直接端点证据。两组完整ABC的X/Y也变了，不能把整组ABC效果差全归于背景。两法在本图使用相同实体团队，不能据负尾归咎某一种调参。','E20')
    p('案例诊断合计7张图、8套不同ABC、24个唯一端点的full_state重放，保存219组、2,066条观测票；这些是已有端点的诊断重放，不增加本轮16,810个真正新增物理团队的计数。')
    p('旧pa4otMbVnkk-05的接缝与自遮挡案例继续保留统计事件，但不把单张RGB中不清楚的两个底角作为“必需结构恢复”的主视觉证据。若人工要继续看，应判断两个局部是否有可见边界或唯一Manhattan推断依据；不因分数改善重新定义范围。当前尚需人工的具体图、位置和判断问题已逐条列在案例表；统计和算法机制已经可复核，不要求先人工裁定才能交付研究结论。','E12/E14')

    h('十三、人员名单与实际可执行替换')
    p('以下名单是ALL描述分区，适合阅读画像；实际事件必须按目标建筑的LOBO标签。名单文件包含62人、历史支持、缺失原因，逐楼外标签另存。TARGET/RAW3的未知者为W019、W022、W026；OPPJOINT另有W014、W018、W034未获联合配置标签。未知者不是“零质量类”。')
    centers=pd.read_csv(SYN/'ADOPTED_CLASS_MEMBERS_AND_PROFILES_ALL.csv')
    rows=[]
    sem={'C1':'困难结构较简化、近共线冗余较多','C2':'常规/中间结构表达','C3':'复杂机会中保留较多转折和短边'}
    for _,r in centers[centers.method.eq('TARGET_DSTRUCT_K3')].sort_values('class_id').iterrows():
        rows.append([r.class_id+'：'+sem[r.class_id],str(int(r.workers)),r.members.replace('|','、')])
    table('TARGET_DSTRUCT K3主名单',['类别含义','人数','成员'],rows,'E16')
    rows=[]
    for cl in ['C1','C2','C3']:
        r=select('E23',method='TARGET_DSTRUCT_K3',class_id=cl)
        rows.append([cl]+[f(r[c],3,True) for c in ['m_hard_turn_n','m_all_short_turn_edges','m_all_near_straight_nodes_5deg','m_all_unmatched_ref_9','m_all_B8L2']])
    table('实际标法的历史画像：相对任务同行的残差',['类','困难转折','总体短边','近共线节点','未匹配参考角','历史Q（后验）'],rows,'E23')
    p('这三类的区别主要在复杂结构机会下保留多少转折、短边以及怎样加入近共线节点，不是给人永久贴上“好/坏”标签。表中值是按原历史层级残差汇总，不是每张图的绝对点数；未匹配参考角同时混合缺漏和定位。Q只作分类后的画像说明，未输入TARGET拟合。C1与C2的困难转折均偏少，但近共线冗余不同；C3的更多转折在一些图恢复必要细部，在另一些图也会带来临界票或连接冲突。')
    p('RAW3较多转折类由原12人扩至17人，新增W011、W012、W013、W033、W036；较多转折组困难残差均值由约+1.796降至+1.356，说明定义确实变宽。其ALL总17人不表示每张图有17人，也不支持从该类取两个互斥七人源块；当前大块高转折纯源反向容量仍未打开。','E04')
    rows=[]
    for method,name in [('V21_DSTRUCT_RAW3_K3','RAW3 K3'),('OPPJOINT_K3','OPPJOINT K3'),('OPPJOINT_K4','OPPJOINT K4'),('RFUSE_GEOM_K2','RFUSE K2')]:
        for _,r in centers[centers.method.eq(method)].sort_values('class_id').iterrows():
            rows.append([name+' '+r.class_id,str(int(r.workers)),r.members.replace('|','、')])
    table('互补方法ALL完整名单',['方法/类','人数','成员'],rows,'E16')
    p('实际展示建议按三个尺度组织：TARGET12/4先体现宽覆盖；TARGET14/7及20/7展示同类B与跨类C的差异及人数背景响应；RAW3 12/8与16/8展示八人过半/半队的扩覆盖，并同时展示较少表达方向的降质与失败。联合画像用OPPJOINT K3做容量更足的八人方向，K4保留慢小类与原六人父子对照。每个图的D/X/Y/Z已列入事件表，可直接回连，不需要从ALL名单临时拼人。')

    h('十四、与当前采用方案相比，保留、修正和新增')
    table('结论变化与证据来源',['类别','本轮处理','由什么回答'],[
      ['保留','v2.1数据/参考/范围；已有Q基线、真K3/K4父子、反向变体；原联合升Q与时间角色；RFUSE反向互补','旧完整结果重读与身份核查，不重复跑曲线'],
      ['修正','OPPJOINT从总主分类改为联合画像主层；TARGET K3成为普通Manual结构主分类','全45比较、原答/稳定性、严格ABC与公平支持'],
      ['修正','旧约+8对+3.9不能单独归因联合结构/时间；第四类不是主升Q方向的必要条件','相同D/A重合为0、连续历史Q梯度、真父事件身份'],
      ['新增','15个v2.1定向变体、支持门槛与原参数精确复现；RAW3作为八人补充','本轮楼外拟合、原答几何和6548严格ABC'],
      ['新增','七八人全部容量漏斗、四主格缺口补齐、真正半队/过半与较大队分开','真实池逐类容量、设计登记、端点状态与逐事件指标'],
      ['新增','同团队Q/QT条件增量、标签容量随机参照、同配比成员残余','复用既有有限团队，未重融合'],
      ['新增','六类机制与第七负尾回查，记录劳动量和Q/IoU不一致并列','原图、原答、完整几何、full_state票账']],'E04/E05/E12/E14/E17')
    p('本轮已足以给出清楚采用方案。未解决的问题有具体边界：类别仍不是永久人型；部分图片的遮挡底角、主顶面/梁底解释需要人工；旧大队部分定位几何缺失；m8来源匹配时间覆盖有限；所有主结果是当前有限人员/图片池下的描述，不能自动推广到新人员或全新建筑。继续扩展应针对这些已发现问题，而不是为名称对称或“所有图片同向”增加实验。')

    h('十五、完整交付与数值回连')
    table('分阶段目录',['目录','内容','入口'],[
      ['stage00_intake','源读取、版本与研究合同、旧采用快照','INTAKE_AND_COMPUTATION_CONTRACT.json'],
      ['stage01_classification','参数、23折名单、原答几何、门槛、真父子与代码','results/DSTRUCT_V21_PARAMETER_AND_RAW_AUDIT_ZH.md'],
      ['stage02_comparison / stage02_capacity','旧全比较、固定团队Q/QT、残余、全漏斗、逐事件与共同支持','FINAL_EVIDENCE_REVIEW_ZH.md；FINAL_COMPETITIVE_N_M_FUNNEL.csv'],
      ['stage03_completed_abc','冻结设计、6548事件、16818端点标量/完整几何、全部769指标','EVENTS_EXACT.csv.gz；EVENT_METRICS.csv.gz；ENDPOINT_GEOMETRY.jsonl.gz'],
      ['stage04_cases','真实全景、原答叠图、局部图、完整full_state与逐票账','独立案例浏览页及每例说明'],
      ['stage05_synthesis','本报告、60配置采用决定、主名单、数值提取与证据索引','ALL60_CURRENT_ADOPTION_DECISIONS.csv；REPORT_TABLE_EXTRACTIONS.csv'],
      ['stage06_reproduction','必要输入、当前原算法、冻结引擎、来源清单','REPRODUCTION_README_ZH.txt；SOURCE_FILE_MANIFEST.csv']],'E01/E17')
    p('所有正文主表均回连实际结果文件和过滤键；报告生成时保存了原始行值，不只保存四舍五入后的表。压缩包检查包括逐成员CRC和内部gzip读到EOF；完整交付再给SHA256清单及可读性回执。旧版本保留，当前新增表以独立目录和文件名保存。')

    h('附录A：证据键与精确路径')
    for key,v in SOURCES.items():p(f"[{key}] {v['description']}：{v['delivery_relative_path']}")
    h('附录B：方法文献（英文原始文献）')
    p('[M1] Ward, Joe H., Jr. (1963). “Hierarchical Grouping to Optimize an Objective Function.” Journal of the American Statistical Association, 58(301), 236–244. DOI: 10.1080/01621459.1963.10500845. 官方PDF：https://www.tandfonline.com/doi/pdf/10.1080/01621459.1963.10500845 （出版方访问权限可能适用）。用于层次聚类的方法来源，不为本项目人员类别提供外部真值。')
    p('[M2] Hubert, Lawrence, and Phipps Arabie (1985). “Comparing Partitions.” Journal of Classification, 2, 193–218. DOI: 10.1007/BF01908075. 官方PDF：https://link.springer.com/content/pdf/10.1007/BF01908075.pdf （出版方访问权限可能适用）。用于分区一致性的解释；ARI不等于分类准确率。')
    p('[M3] Varma, Sudhir, and Richard Simon (2006). “Bias in Error Estimation When Using Cross-validation for Model Selection.” BMC Bioinformatics, 7, 91. DOI: 10.1186/1471-2105-7-91. 作者所在机构开放PDF：https://linus.nci.nih.gov/techreport/Varma-Simon-CrossValid.pdf 。用于说明方法选择与独立预测误差评估必须区分；本报告没有声称完成外部预测验证。')


def font(run,size=None,bold=None):
    run.font.name='Noto Sans CJK SC';run.font.color.rgb=RGBColor(0,0,0)
    if size:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    rp=run._element.get_or_add_rPr();rf=rp.rFonts
    if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
    for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'Noto Sans CJK SC')


def render_docx():
    d=Document();sec=d.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    for name in ['Normal','Title','Heading 1','Heading 2','Caption','Header','Footer']:
        style=d.styles[name];style.font.name='Noto Sans CJK SC';style.font.color.rgb=RGBColor(0,0,0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC')
    for border in d.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    normal=d.styles['Normal'];normal.font.size=Pt(10);normal.paragraph_format.line_spacing=1.28;normal.paragraph_format.space_after=Pt(5)
    for name,size in [('Title',21),('Heading 1',14),('Heading 2',11.5)]:
        st=d.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6)
    hp=sec.header.paragraphs[0];r=hp.add_run('HOHONET v2.1  |  普通 Manual 人员分类独立复核');font(r,8)
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=fp.add_run('2026-10-29  ·  ');font(r,8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
    for block in BLOCKS:
        if block['kind']=='h':
            pp=d.add_paragraph(style='Title' if block['level']==0 else 'Heading '+str(block['level']));font(pp.add_run(block['text']))
        elif block['kind']=='p':
            pp=d.add_paragraph()
            if block['text'].startswith('[E'):
                pp.paragraph_format.line_spacing=1.1
                pp.paragraph_format.space_after=Pt(4)
                pp.paragraph_format.keep_together=True
                font(pp.add_run(block['text']),8.5)
            elif block['text'].startswith('[M'):
                pp.paragraph_format.line_spacing=1.1
                pp.paragraph_format.space_after=Pt(6)
                pp.paragraph_format.keep_together=True
                font(pp.add_run(block['text']),9)
            else:font(pp.add_run(block['text']))
            if block.get('source'):font(pp.add_run(' ['+block['source']+']'),8)
        elif block['kind']=='table':
            pp=d.add_paragraph();pp.paragraph_format.keep_with_next=True;font(pp.add_run(block['title']),10,True)
            t=d.add_table(rows=1,cols=len(block['headers']));t.style='Table Grid';t.autofit=True
            if block['headers'][-1]=='成员':
                t.autofit=False
                widths=[4.2,1.1,11.7]
                for col,width in zip(t.columns,widths):col.width=Cm(width)
            head=t.rows[0]
            if block['headers'][-1]=='成员':
                for cell,width in zip(head.cells,[4.2,1.1,11.7]):cell.width=Cm(width)
            trpr=head._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            for cell,txt in zip(head.cells,block['headers']):
                pp=cell.paragraphs[0];pp.paragraph_format.space_after=Pt(2);font(pp.add_run(str(txt)),8.5,True)
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');cell._tc.get_or_add_tcPr().append(sh)
            for vals in block['rows']:
                rr=t.add_row();trpr=rr._tr.get_or_add_trPr();el=OxmlElement('w:cantSplit');trpr.append(el)
                if block['headers'][-1]=='成员':
                    for cell,width in zip(rr.cells,[4.2,1.1,11.7]):cell.width=Cm(width)
                for c,txt in zip(rr.cells,vals):
                    pp=c.paragraphs[0];pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(2);pp.paragraph_format.line_spacing=1.12;font(pp.add_run(str(txt)),8.4)
            pp=d.add_paragraph();font(pp.add_run('证据：['+block['source']+']。完整精度、其他方向及逐图/楼分母见相应CSV。'),8)
        elif block['kind']=='image':
            d.add_picture(block['path'],width=Cm(17))
            pp=d.add_paragraph(style='Caption');font(pp.add_run(block['caption']),8.5)
    d.core_properties.title='HOHONET v2.1 人员分类独立复核'
    d.core_properties.subject='普通Manual主方法、v2.1定向调参、有限配比、严格ABC与机制'
    d.core_properties.author='HOHONET research audit'
    dest=SYN/'HOHONET_v21_personnel_research_report_ZH.docx';d.save(dest)
    return dest


def render_html():
    parts=['<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>HOHONET v2.1 人员分类独立复核</title><style>body{font-family:system-ui,"Noto Sans CJK SC",sans-serif;color:#111;background:white;max-width:1120px;margin:36px auto;padding:0 26px;line-height:1.85}h1{font-size:30px}h2{font-size:23px;margin-top:38px;border-bottom:1px solid #888;padding-bottom:8px}p{margin:12px 0}table{border-collapse:collapse;width:100%;font-size:14px;line-height:1.65}td,th{border:1px solid #aaa;padding:8px 10px;vertical-align:top}th{background:#eee;text-align:left}figure{margin:22px 0}img{width:100%;height:auto}small,figcaption{font-size:13px;color:#333}.tablewrap{overflow:auto}.evidence{font-size:12px;overflow-wrap:anywhere}@media print{body{max-width:none;font-size:10pt}h2{break-after:avoid}tr{break-inside:avoid}}</style><body>']
    for b in BLOCKS:
        if b['kind']=='h':parts.append(f'<h{1 if b["level"]==0 else b["level"]+1}>{html.escape(b["text"])}</h{1 if b["level"]==0 else b["level"]+1}>')
        elif b['kind']=='p':parts.append('<p>'+html.escape(b['text'])+((' <small>['+html.escape(b['source'])+']</small>') if b.get('source') else '')+'</p>')
        elif b['kind']=='table':
            parts.append('<h3>'+html.escape(b['title'])+'</h3><div class="tablewrap"><table><thead><tr>'+''.join('<th>'+html.escape(str(t))+'</th>' for t in b['headers'])+'</tr></thead><tbody>')
            for row in b['rows']:parts.append('<tr>'+''.join('<td>'+html.escape(str(t))+'</td>' for t in row)+'</tr>')
            parts.append('</tbody></table></div><p class="evidence">证据：['+html.escape(b['source'])+']</p>')
        elif b['kind']=='image':
            data=base64.b64encode(Path(b['path']).read_bytes()).decode();parts.append('<figure><img src="data:image/png;base64,'+data+'" alt="'+html.escape(b['caption'],quote=True)+'"><figcaption>'+html.escape(b['caption'])+'</figcaption></figure>')
    parts.append('</body></html>');(SYN/'HOHONET_v21_personnel_research_report_ZH.html').write_text(''.join(parts))


if __name__=='__main__':
    build_content()
    (SYN/'REPORT_BLOCK_SOURCE.json').write_text(json.dumps(BLOCKS,ensure_ascii=False,indent=2))
    (SYN/'REPORT_EVIDENCE_INDEX.json').write_text(json.dumps(SOURCES,ensure_ascii=False,indent=2))
    pd.DataFrame(EXTRACTS).to_csv(SYN/'REPORT_TABLE_EXTRACTIONS.csv',index=False,encoding='utf-8-sig')
    path=render_docx();render_html()
    print(json.dumps({'docx':str(path),'blocks':len(BLOCKS),'tables':sum(x['kind']=='table' for x in BLOCKS),'source_rows':len(EXTRACTS)},ensure_ascii=False))
