"""把已完成的研究验证转为中文报告、研究图与客观Word；不再拟合模型。"""
import json
import sys
from pathlib import Path

ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'analysis_results/research_validation_20260909_v2'


def fmt(value,digits=3):
    if value is None or str(value)=='nan': return '—'
    return f'{value:.{digits}f}' if isinstance(value,(float,int)) else str(value)


def markdown_table(headers,rows):
    return '\n'.join(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']+
                     ['| '+' | '.join(str(x) for x in row)+' |' for row in rows])


def build_content():
    import numpy as np
    import pandas as pd
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    if str(ROOT) not in sys.path: sys.path.insert(0,str(ROOT))
    from tools.thesis_main.analysis.validate_building_stages_20260909 import OLD

    figs=OUT/'figures'; figs.mkdir(parents=True,exist_ok=True)
    font_manager.fontManager.addfont('C:/Windows/Fonts/msyh.ttc')
    plt.rcParams.update({'font.family':'Microsoft YaHei','axes.unicode_minus':False,'font.size':10,
                         'axes.spines.top':False,'axes.spines.right':False,'savefig.facecolor':'white'})
    inv=pd.read_csv(OLD/'comparison/image_support_comparison.csv').set_index('image_id')
    on=pd.read_csv(OUT/'stages/image_stage_onsets.csv')
    curves=pd.read_csv(OUT/'stages/stage_curves.csv.gz')
    summary=pd.read_csv(OUT/'prediction/prediction_summary.csv')
    buildings=pd.read_csv(OUT/'prediction/per_building_scores.csv')
    features=pd.read_csv(OUT/'features/image_feature_audit.csv')
    feature_plan=json.loads((OUT/'features/FEATURE_PLAN.json').read_text(encoding='utf8'))
    cohort={'with_workers':'含W19/W26','without_workers':'排除W19/W26'}
    label={'identified':'起点可识别','not_reached':'观察范围内未达到','unknown':'起点仍有不确定'}
    dense=inv[inv.common_n>=16]
    eligible=dense[dense.building_id.map(dense.groupby('building_id').size())>=2]
    main=on[(on.kind=='anchor_stage')&(on.horizon==20)]
    main50=main[main.image_id.isin(eligible.index)]
    main50.groupby(['building_id','mode','config','status']).size().unstack(fill_value=0).to_csv(OUT/'stages/building_stage_summary.csv')
    qsummary=main.groupby(['mode','config','status']).size().unstack(fill_value=0).reset_index()
    qsummary.to_csv(OUT/'stages/q_sensitivity_summary.csv',index=False)
    # 阈值敏感性只读取已保存曲线，不据结果选择主阈值。
    rate_rows=[]
    c=curves[(curves.kind=='anchor_stage')&(curves.horizon==20)]
    for key,g in c.groupby(['image_id','mode','config']):
        g=g.sort_values('k')
        for rate in (.7,.8,.9):
            def first(values):
                ix=np.flatnonzero(np.asarray(values)[1:]>=rate-1e-12)
                return int(ix[0]+2) if len(ix) else None
            lo,hi=first(g.upper),first(g.lower)
            status='identified' if lo is not None and lo==hi else 'not_reached' if lo is None else 'unknown'
            rate_rows.append(dict(zip(['image_id','mode','config'],key))|dict(rate=rate,possible=lo,conservative=hi,status=status))
    pd.DataFrame(rate_rows).to_csv(OUT/'stages/rate_sensitivity.csv',index=False)
    # 对既有逐楼贡献的条件重抽样，不冒充新楼/新人验证。
    gain_rows=[]; rng=np.random.default_rng(20260909)
    for key,g in buildings[(buildings.experiment=='high_split')&(buildings.horizon==20)&(buildings.design=='balanced')&
                            buildings.model.isin(['outside_risk','within_risk','outside_shared','outside_compressed','outside_refined'])].groupby(['mode','config','fraction','model']):
        ix=rng.integers(0,len(g),size=(5000,len(g)))
        lo=g.gain_over_base_lower.to_numpy(); hi=g.gain_over_base_upper.to_numpy()
        gain_rows.append(dict(zip(['mode','config','fraction','model'],key))|dict(buildings=len(g),
            gain_lower=lo.mean(),gain_upper=hi.mean(),conditional_resample_lower=np.quantile(lo[ix].mean(1),.025),
            conditional_resample_upper=np.quantile(hi[ix].mean(1),.975)))
    pd.DataFrame(gain_rows).to_csv(OUT/'prediction/building_resample_sensitivity.csv',index=False)

    colors=['#9EA8B3','#E8C078','#5588B4']
    fig,ax=plt.subplots(figsize=(8.8,3.2))
    counts=[int(inv.common_n.le(2).sum()),int(inv.common_n.between(3,7).sum()),int(inv.common_n.ge(16).sum())]
    bars=ax.barh(['N=0–2：保留库存','N=3–7：长期预测待验证','N=19–24：高人数历史'],counts,color=colors,height=.6)
    for b,v in zip(bars,counts): ax.text(v+1,b.get_y()+b.get_height()/2,str(v),va='center')
    ax.set(xlabel='图片数（合计214图、22个building）',xlim=(0,115)); ax.invert_yaxis(); fig.tight_layout()
    fig.savefig(figs/'coverage.png',dpi=240); plt.close(fig)

    unb=on[(on.building_id=='uNb9QFRL6hY')&(on.config=='q_0.950')&(on.kind=='anchor_stage')&(on.horizon==23)]
    chosen=['1096f195','d02f87bb','bcce4f23','07a43087','3b9e548b','978d7a8e']
    fig,axes=plt.subplots(2,3,figsize=(10.5,6.2),sharex=True,sharey=True)
    for ax,short in zip(axes.flat,chosen):
        for mode,color in [('with_workers','#24699B'),('without_workers','#C86B39')]:
            g=curves[(curves.image_id.str.contains(short))&(curves.config=='q_0.950')&(curves.kind=='anchor_stage')&(curves.horizon==23)&(curves['mode']==mode)].sort_values('k')
            ax.fill_between(g.k,g.lower,g.upper,color=color,alpha=.15)
            ax.plot(g.k,g.lower,color=color,lw=1.8,label=cohort[mode]); ax.plot(g.k,g.upper,color=color,ls=':',lw=1)
        ax.axhline(.8,color='#555',ls='--',lw=.8); ax.set_title(short); ax.set(ylim=(0,1.04),xticks=[2,6,10,14,18],xlim=(2,18))
    axes[0,0].legend(fontsize=8,loc='upper left'); fig.supxlabel('候选起点 k；每条顺序持续检查到 H=23')
    fig.supylabel('持续阶段达标率的未知状态边界'); fig.tight_layout()
    fig.savefig(figs/'unb_stage_curves.png',dpi=240); plt.close(fig)

    config_names=['q_0.975','q_0.950','q_0.925','q_0.900','q_0.850','q_0.800','ospa30_t6','hac_q_0.950']
    labels=['q=.975','q=.950','q=.925','q=.900','q=.850','q=.800','OSPA≤6°','凝聚q=.950']
    fig,axes=plt.subplots(1,2,figsize=(10.5,4.2),sharey=True)
    for ax,mode in zip(axes,['with_workers','without_workers']):
        g=qsummary[qsummary['mode']==mode].set_index('config').reindex(config_names); left=np.zeros(len(g))
        for state,color in [('identified','#477FA5'),('not_reached','#D99655'),('unknown','#B9C1C9')]:
            ax.barh(labels,g[state],left=left,color=color,label=label[state]); left+=g[state].to_numpy()
        ax.set(title=cohort[mode],xlabel='图片数（H=20，共54图）',xlim=(0,54)); ax.invert_yaxis()
    axes[0].legend(fontsize=8,loc='lower left',bbox_to_anchor=(0,-.26),ncol=3)
    fig.tight_layout(); fig.savefig(figs/'q_sensitivity.png',dpi=240,bbox_inches='tight'); plt.close(fig)

    psel=summary[(summary.experiment=='high_split')&(summary.horizon==20)&(summary.fraction==.6)&(summary.design=='balanced')&(summary.config=='q_0.950')]
    selected_models=['building_mean','within_risk','within_shared','outside_mean','outside_risk','outside_compressed','outside_refined','outside_shared']
    model_labels=['同楼历史均值','同楼＋风险近邻','同楼＋共享层近邻','楼外历史均值','楼外风险近邻','楼外压缩层近邻','楼外精炼层近邻','楼外共享层近邻']
    fig,axes=plt.subplots(1,2,figsize=(10.5,4.3),sharey=True)
    for ax,mode in zip(axes,['with_workers','without_workers']):
        g=psel[psel['mode']==mode].set_index('model').reindex(selected_models)
        for y,row in enumerate(g.itertuples()):
            ax.plot([row.curve_error_lower,row.curve_error_upper],[y,y],lw=5,color='#3D7FA3',solid_capstyle='round')
        ax.set(yticks=range(len(g)),yticklabels=model_labels,title=cohort[mode],xlabel='曲线平均绝对误差的可行范围',xlim=(0,.65)); ax.invert_yaxis(); ax.grid(axis='x',alpha=.2)
    fig.tight_layout(); fig.savefig(figs/'prediction_bounds.png',dpi=240); plt.close(fig)

    def get_pred(mode,model,config='q_0.950'):
        return summary[(summary.experiment=='high_split')&(summary.horizon==20)&(summary.fraction==.6)&(summary.design=='balanced')&
                       (summary.config==config)&(summary['mode']==mode)&(summary.model==model)].iloc[0]
    count_with=get_pred('with_workers','building_mean')
    count_without=get_pred('without_workers','building_mean')
    def p(text): return {'type':'p','text':text}
    def table(headers,rows): return {'type':'table','headers':headers,'rows':[[str(x) for x in r] for r in rows]}
    def fig(name,caption,width=6.4): return {'type':'figure','path':str(OUT/name),'caption':caption,'width':width}
    pages=[]
    scope_rows=[['历史库存','214图 / 22楼 / 26人','2501份canonical；2513份版本'],
        ['无辅助可计算','196图；1887 / 1843响应','包含 / 排除W19、W26'],['高人数历史','55图 / 15楼','共同人数19–24'],
        ['同楼留图预测','50图 / 10楼','约50% / 60%作历史源'],['低人数待验证','98图 / 21楼','N=3–7，已生成长期预测'],
        ['库存保留','61图','18图N=0；20图N=1；23图N=2']]
    pages.append({'title':'相似场景的标注稳定阶段与人员组合','subtitle':'历史数据验证、图像特征对照与新增标注设计｜2026年9月9日','blocks':[
        p('研究对象是标注分布在有限人数范围内进入持续稳定阶段的时间，以及该阶段能否由同building历史图和图像特征预测。稳定多个簇属于允许的结果；分布稳定、参考偏差和人员类型复现分别评价。'),
        table(['数据层级','覆盖','统计口径'],scope_rows),
        fig('figures/coverage.png','图1 共同人员预算下的图像覆盖。低人数图没有被当作长期不稳定样本。',6.2),
        p('所有回放均复用真实人员的独立无辅助作答。200条人员顺序和重复留图组合不增加独立人员、图片或building数量。')]})
    pages.append({'title':'本轮检验问题与稳定阶段定义','blocks':[
        table(['问题','本轮可观测量','评价对象'],[
            ['人数增加后是否出现阶段','持续稳定达标率、起点范围；代表参考偏差曲线','每图真实人员前缀'],
            ['阶段能否跨图预测','同楼源、楼外源、图像特征的留图误差','整张目标图留出'],
            ['人员组合能否复用','训练外分档、同图不同人组合、代表偏差与分歧','固定2人或4人预算']]),
        p('每簇必须满足有效点数一致。有效点数不同的两份作答一定分属不同簇，即使坐标接近。几何组至少有两位不同人员支持时才成为受支持簇；单例保留为低支持组，继续参与份额及成员关系计算。'),
        p('在一条人员顺序中，对候选前缀k与其后每个j直到观察终点H进行比较：旧人员同簇关系改变的人员对比例≤0.10；各簇份额总变差TV≤0.10；没有簇相对k从不足2人跨到至少2人；k时已有至少一个受支持簇。'),
        p('持续阶段要求同一条顺序从k到H−5的每个起点都满足上述对照。合法候选域为k≥2，之后至少有5位人员；H=20时评价2≤k≤15，H=23时评价2≤k≤18。k=1仅保留轨迹，不进入起点提取。'),
        table(['状态','定义'],[['stable','所有必要比较满足条件；允许稳定单簇和多簇'],['changing','至少一个可评价比较已经违反条件'],['unknown','没有已知违反，但必要分区或支持仍不足']]),
        p('L=stable/200，U=(stable+unknown)/200。起点以80%顺序进入持续阶段为阈值，分别求U与L过线人数；同一有限值才记为起点可识别。L/U是未知状态边界，不是置信区间；“可识别”也不意味着排除了200次回放的蒙特卡洛波动。'),
        p('未在H内达到条件，不代表永久不收敛；分布稳定也不代表正确或理论质量上限。')]})
    b_rows=[]
    for b,g in dense.groupby('building_id'):
        if len(g)>=2:
            b_rows.append([b,len(g),f'{int(g.common_n.min())}–{int(g.common_n.max())}', '同楼＋楼外'])
    b_rows.append(['5个单图building',5,'19–24','楼外单列'])
    headline=[]
    for mode in cohort:
        v=main50[(main50['mode']==mode)&(main50.config=='q_0.950')].status.value_counts()
        headline.append(f"{cohort[mode]}：50图中起点可识别{v.get('identified',0)}图、观察内未达到{v.get('not_reached',0)}图、起点不确定{v.get('unknown',0)}图。")
    pages.append({'title':'所有具备高人数历史的building','blocks':[
        p('高人数入口由不同人员数预先确定，不按是否稳定选图。15个building共有55图；其中10楼各有至少2图，共50图，可以留出同楼新图。其余5楼各只有一张高人数图。'),
        table(['building','高人数图','共同人数范围','预测评价'],b_rows),
        p('单图building：VFuaQ6m2Qom、X7HyMhZNoso、e9zR4mvMWw7、pRbA3pwrgk9、x8F5xyUWy9e。'),
        p('主对照H=20。VFuaQ6m2Qom仅19人，纳入55图H=19的楼外检验；H=22/23/24补充结果另列，不混算不同观察终点。'),
        p('q=.95、H=20、持续阶段主定义：'+' '.join(headline)),
        p('同building只定义当前分组，不代表所有图属于同一物理房间。房间身份未在本轮验证；不同图的稳定阶段差异不能自动归因于房间类别。')]})
    unb_rows=[]
    for short in chosen:
        values=[]
        for mode in cohort:
            r=unb[(unb.image_id.str.contains(short))&(unb['mode']==mode)].iloc[0]
            values.append(str(int(r.identified)) if pd.notna(r.identified) else label[r.status])
        unb_rows.append([short,*values])
    pages.append({'title':'uNb：同楼图像存在不同的稳定阶段','blocks':[
        p('uNb共有12张高人数图。本页统一H=23、q=.95、两人支持和0.10变化容差。约7人、11–14人、18人的起点同时出现；还存在观察范围内未达到和起点不确定的图。'),
        fig('figures/unb_stage_curves.png','图2 六张uNb图的持续阶段曲线；阴影表示未知状态范围，横虚线为80%。',6.3),
        table(['图像短ID','含W19/W26','排除W19/W26'],unb_rows),
        p('07a43087在k=18时有182/200与189/200条顺序达标；这些达标顺序全部包含至少两个各有≥2人支持的簇。该图的多簇结构没有被判成稳定失败。'),
        p('新定义要求同一顺序持续至H；旧h5摘要与这些人数不能混算。')]})
    qrows=[]
    for mode in cohort:
        for config in ['q_0.975','q_0.950','q_0.925','q_0.900','ospa30_t6','hac_q_0.950']:
            r=qsummary[(qsummary['mode']==mode)&(qsummary.config==config)].iloc[0]
            qrows.append([cohort[mode],config,int(r.identified),int(r.not_reached),int(r.unknown)])
    pages.append({'title':'分簇方法与q阈值的敏感性','blocks':[
        fig('figures/q_sensitivity.png','图3 H=20的54图；相同数据、相同人员顺序与稳定标准，仅改变几何配置。',6.4),
        p('v5一系保留逐次最大完全子图产生的全部候选分区，仅在分区唯一时给出簇归属。凝聚complete-link补充按连续几何距离合并；仍限制同点数且簇内任意两份距离达标，但相等距离可能受固定输入顺序影响。'),
        p('q=.95不是95%置信度。按现有公式，它对应上下边界平均绝对差不超过25.6像素（1024×512画布），以及墙事件双向最近横向距离均值不超过25.6像素，即约9°。这些是平均通道误差，不是所有角点都在9°以内；还要求现有配对/循环对应可用。'),
        table(['q','边界平均差上限（像素）','墙事件平均横向角距'],[['.975','12.8','4.5°'],['.950','25.6','9°'],['.925','38.4','13.5°'],['.900','51.2','18°']]),
        p('当前没有独立人工“应同簇／应异簇”标定集，因此本轮没有确定最优q。q=.95作为固定对照保留，完整网格和70%/80%/90%阶段达标率敏感性均已落盘。算法减少不确定分区，不等于证明某个分区是真实语义。'),
        p('人工复核的删点、补点和响应暂排已沿用。撤回两份补点的敏感性中，部分图由“观察内未达到”转为“仍不确定”；两图没有因此获得新的可识别起点。')]})
    dsg=pd.read_csv(OUT/'prediction/split_manifest.csv')
    split_rows=[]
    for b,g in dsg[dsg.design=='balanced'].groupby('building_id'):
        a=g[g.fraction==.5].iloc[0]; z=g[g.fraction==.6].iloc[0]
        split_rows.append([b,f'{a.source_n}→{a.target_n}',f'{z.source_n}→{z.target_n}',f'{len(g[g.fraction==.5])} / {len(g[g.fraction==.6])}'])
    pages.append({'title':'留图预测的分配和评价','blocks':[
        p('先把每楼图像划为历史源和互补目标，再预测目标的持续阶段。目标图完整标注只用于评价。分配用标注前图像风险的楼内分位段、共同人数做边际平衡，同时保留未筛平衡条件的对照。每类最多保留200种可行划分，小规模时全部枚举。'),
        table(['building','约50%历史→目标','约60%历史→目标','平衡划分数'],split_rows),
        p('6楼40图可在每折留出至少2个目标；4个小楼的10图退化为每折1个目标，已显式列出。不能声称全部50图都具有多目标留图。两个人员口径复用同一划分。'),
        p('每个留出目标同时比较：同楼历史平均曲线；楼外所有合格历史的平均曲线；按固定图像特征选至多3个近邻并以距离倒数加权；同楼均值与楼外近邻各占一半的预设组合。楼外预测排除目标building的所有人类结果。'),
        p('楼外平均使用更大的候选历史池，是跨楼全池均值基线；它不是与楼内源张数相同的单次随机样本。近邻预测最多取3张。源张数、候选池覆盖及其差异不归因于building本身。'),
        p('主指标在k=2至H−5计算曲线误差边界，先平均同一目标的重复划分，再按building等权。条件人数误差仅对目标可识别、来源具有可识别起点的可预测部分计算；必须同时报告图数。616884条预测记录不等于616884次独立实验。')]})
    pred_rows=[]
    for model,name in zip(selected_models,model_labels):
        a=get_pred('with_workers',model); b=get_pred('without_workers',model)
        pred_rows.append([name,f'[{a.curve_error_lower:.3f}, {a.curve_error_upper:.3f}]',f'[{b.curve_error_lower:.3f}, {b.curve_error_upper:.3f}]'])
    pages.append({'title':'预测结果：图像风险和特征层对照','blocks':[
        p('下图为10楼50图、H=20、q=.95、约60%历史源的平衡划分结果。线段是平均绝对误差的可行范围，越靠左误差越小；它不是置信区间。'),
        fig('figures/prediction_bounds.png','图4 全部50个目标均计入曲线评价；未知状态不会因为误差下界为零被视作准确预测。',6.4),
        p('同楼均值的误差范围为[0.220, 0.560] / [0.254, 0.574]；楼外风险近邻为[0.111, 0.390] / [0.156, 0.404]。q主分析中，楼外风险近邻相对同楼均值的配对改善范围仍跨零：[-0.108, 0.392] / [-0.096, 0.362]。'),
        p('OSPA对照中，相同风险预测的楼等权平均改善范围为[0.005, 0.357] / [0.012, 0.363]，在固定历史数据及未知边界下方向为正。对10栋楼贡献重抽样后仍需考虑楼际变动；该范围不是跨新楼或新人的泛化置信保证。'),
        p(f'同楼均值的条件人数MAE为{count_with.count_error:.3f} / {count_without.count_error:.3f}人，分别只覆盖{int(count_with.count_buildings)}楼{int(count_with.count_targets)}图和{int(count_without.count_buildings)}楼{int(count_without.count_targets)}图；有计数评分的目标出现为{int(count_with.count_appearances)}/{int(count_with.target_appearances)}与{int(count_without.count_appearances)}/{int(count_without.target_appearances)}。各模型可预测图与划分不同，未据此单独排名。')]})
    shape_rows=[]
    names={'encoder_2':'编码器第2阶段','encoder_4':'编码器第4阶段','compressed':'水平压缩后','refined':'水平精炼后','shared':'共享潜在特征','legacy_mean_phase0':'旧均值汇聚形式（单相位）'}
    for name in names:
        shape_rows.append([names[name],feature_plan['feature_shapes'][name][1], '单相位均值＋L2' if name=='legacy_mean_phase0' else '空间均值与标准差；四相位平均'])
    pages.append({'title':'HoHoNet特征与低人数图的纳入','blocks':[
        p('214图已在统一CPU环境、ep300布局模型和同一输入选取规则下重新提取特征。各层来自同一次前向计算；模型没有使用本次人类稳定标签训练。d_model_feat是共享特征到冻结参考空间的近邻距离，不是已经标定的人类难度或失败概率。'),
        table(['特征位置','维数','汇聚方式'],shape_rows),
        p('主比较覆盖编码器中间/末端、水平压缩、水平精炼及共享特征。没有一种原始层向量在全部距离配置、人员口径和评价覆盖下均优于风险距离。压缩/精炼层的部分误差边界低于共享层，层选择仍属探索。'),
        p('旧d_t为归一化单相位共享均值到Calibration_manual池的第10近邻距离；新分数用四相位均值/标准差及固定训练白化空间。旧汇总risk仅70图有值、高人数仅4图，且不能直接确认为d_t。本轮比较旧均值向量形式，未完成两版分数的等覆盖验证。'),
        table(['混合图集的组成','数量','长期目标状态'],[['高人数留图','50图 / 10楼','可以历史评价'],['N=3–7','98图 / 21楼','已预测H=20曲线，等待补采验证'],['其中有同楼高人数源','78图','同楼及楼外预测并列'],['其中无同楼高人数源','20图','楼外预测单列'],['N≤2','61图','保留库存，不判长期稳定']]),
        p('低人数图的前瞻预测保留在逐图表中，不能与高人数图合并成一个长期准确率。此前uNb的N=5短窗实验只验证短窗状态，不检验8人后至20余人的持续阶段。')]})
    pages.append({'title':'人员差异：连续得分与离散类别','blocks':[
        p('无辅助可计算响应1887条中，1484条、167图具备本轮可用参考；403条参考未定，只进入分歧分析。已确认人工参考优先，其他可评价图暂按GT；偏离参考不自动等于粗心或错误范围选择。'),
        fig('workers/worker_validation.png','图5 人员效应的整楼留出预测和互斥building分半复现；OSPA30/60并列。',6.4),
        table(['当前20人','连续人员效应','中位数2档','Ward3档'],[
            ['楼等权MSE改善（30 / 60）','3.38% / 0.067%','0.42% / −0.76%','1.51% / −1.28%'],
            ['互斥分半复现中位数','排序相关0.598 / 0.642','ARI 0.113 / 0.113','ARI 0.104 / 0.091']]),
        p('人员得分由图像效应调整后的参考偏差得到；每个目标building的人员档位均在其他building重新拟合。A表示训练参考偏差较低，B较高；三档时C最高。全量人员名单只作解释，不用于留楼验证标签。'),
        p('连续排序有一定复现，固定两档或三档的成员复现较弱。指定切成几档不是天然存在几类人群的证据。删一楼和互斥分半得到的稳定度不同，不能用高度重叠训练集的一致率替代独立复现。')]})
    pages.append({'title':'按人采集后的真实人员组合','blocks':[
        p('每位人员对同一图的一份canonical作答可在后处理中进入不同配比；同一人不能复制为两个人。AABC必须有两位不同A及各一位B、C的真实独立作答。组合不包含相互讨论或看到他人结果后的修改效应。'),
        fig('workers/worker_combinations.png','图6 当前20人，按目标楼外资料分A/B；固定四人预算比较五种配比。',6.4),
        table(['四人配比','AAAA','AAAB','AABB','ABBB','BBBB'],[
            ['代表参考偏差OSPA30','3.98','4.24','4.67','4.97','5.14'],
            ['代表参考偏差OSPA60','5.36','5.88','6.71','7.36','7.72'],
            ['成对分歧OSPA30','7.44','7.68','7.86','7.96','7.99'],
            ['成对分歧OSPA60','11.63','12.24','12.54','12.53','12.23']]),
        p('单位为度，先图后楼等权。参考指标共同41图/12楼；分歧共同55图/15楼。代表按组合内部距离选真实medoid，参考答案不参与选择；并列代表均权计分。'),
        p('两种参考偏差均有11/12楼AAAA低于BBBB，B6ByNegPMKs方向相反。当前20人的Ward三档最低偏差档在各留楼折均只有W2，AABC可组合图数为0。组合可计算、类别可复现和某类在20人下收敛是三个不同条件。')]})
    quality=pd.read_csv(OUT/'workers/quality_prefix_summary.csv')
    qr=[]
    for co in ['all26','current20']:
        for metric in ['ospa30','ospa60']:
            g=quality[(quality.cohort==co)&(quality.metric==metric)].set_index('k')
            qr.append([co,metric,f'{g.loc[1,"building_equal_mean"]:.3f}',f'{g.loc[8,"building_equal_mean"]:.3f}',f'{g.loc[20,"building_equal_mean"]:.3f}'])
    pages.append({'title':'人数与参考质量：8人后平均变化较小','blocks':[
        p('在参考可评价且有至少20位真实人员的图上，沿相同200条全局人员顺序，分别用k=1、2、4、6、8、12、16、20人的组内距离选medoid，再评价该真实代表作答的参考偏差。全26人口径固定41图/12楼，当前20人固定40图/12楼。'),
        fig('workers/quality_prefix.png','图7 各人数使用相同图集；曲线描述历史人员抽取下的代表参考偏差。',6.4),
        table(['人员范围','指标','k=1','k=8','k=20'],qr),
        p('8→20人的楼等权变化：全26人OSPA30改善0.012°、OSPA60变差0.062°；当前20人分别改善0.016°、变差0.102°。均值变化较小与逐图改善、恶化并存。'),
        p('这说明给定参考与medoid汇聚方法下，历史平均边际变化已较小；不能证明理论质量上限，不能证明少数解释已完整覆盖，也不能据此认定停止新增受支持簇。当前没有把这条质量曲线替代持续阶段标签。')]})
    pages.append({'title':'新增标注的验证设计与证据范围','blocks':[
        table(['研究命题','已有观测','尚需新增数据检验'],[
            ['有限范围内的阶段差异','同楼存在早、晚、未达到及不确定图；稳定多簇实例','更长人数和新顺序是否保持阶段'],
            ['同楼与图像特征预测','完成50图/10楼留图；风险近邻有参数相关的误差改善','新目标图和新人批次的预测误差及覆盖'],
            ['人员组合复用','真实独立作答可形成多种配比，代表偏差有历史差异','新人员档位及比例效果是否复现'],
            ['质量边际变化','固定图集的8→20平均变化较小','其他汇聚方法、独立参考与更大人数']]),
        p('新增数据按人员独立采集一次。每个building保留多张历史源与多张目标，源和目标两侧都具有高人数重复标注。人员校准图与预测目标图分开；不以目标图结果定义人员档位。'),
        p('在查看目标图完整标注前保存源图、特征版本、分簇阈值、稳定容差、起点/曲线预测和未知范围。补采N=3–7的图可以直接检验本次已保存的H=20预测。人数预算应覆盖预测起点及其后至少5人的观察，而不是把20预设为所有图的停止点。'),
        p('正式评价分别统计可识别目标、观察内未达到、未知和未达到目标观察人数的情况。预测是否有效由预先固定的同楼、楼外和特征基线共同判断，不能只挑预测有利的building、q、人员范围或可识别子集。'),
        p('研究范围以当前可观测量为限：图像稳定阶段与参考质量分别度量；人员档位采用可检验的操作定义；不将“稳定”替换为“正确”，不将“偏差高”替换为“粗心”。')]})
    pages.append({'title':'数据追溯、验证与方法来源','blocks':[
        table(['层级','来源与验证'],[
            ['原始响应','18份Label Studio导出；2501个canonical及2513条版本谱系'],
            ['人工处理计算视图','reviewed：8份删点、2份补点、5份逐份排除；未经确认的奇数点不自动修补'],
            ['几何分簇','revised成对几何与前缀分区；hard effective point count；保留非唯一/不可评价'],
            ['模型','HoHoNet布局ep300；统一214图特征；冻结参考白化空间1647图'],
            ['阶段独立复算','4图×2几何配置×2人员范围×200顺序；48000节点、240聚合节点一致'],
            ['人员组合核对','6179条闭式组合核对；259200条质量前缀排列节点'],
            ['正式边界','没有改变原始标注、正式方法合同、人员资格、派发或停止规则']]),
        p('旧两批特征缓存不能直接混用。87张重叠历史图的向量原本不同；本轮从图像统一重提取。与Stage3的144张重叠图，共享向量最大绝对差约0.000304，重算风险最大差约0.0568；CPU与旧GPU缓存的数值环境不相同，完整差异保留在审计表。'),
        p('HoHoNet的共享潜在水平特征原本用于全景布局等视觉任务；本轮把冻结表示及其参考空间距离作为人类标注阶段的候选预测变量。这项用途由本轮留图实验检验，不能从网络名称或论文的视觉性能直接推出。'),
        {'type':'link','text':'[1] Sun 等，HoHoNet: 360 Indoor Holistic Understanding With Latent Horizontal Features，CVPR 2021。','url':'https://openaccess.thecvf.com/content/CVPR2021/html/Sun_HoHoNet_360_Indoor_Holistic_Understanding_With_Latent_Horizontal_Features_CVPR_2021_paper.html'},
        p('完整逐图数据、所有参数、预测覆盖、人员成员与组合表以及复现命令见同目录研究验证报告及stages、features、prediction、workers子目录。本文为客观结果整理；所有分析均为历史探索，尚无本轮新增真人标注的验证结果。')]})
    payload=dict(title=pages[0]['title'],date='2026-09-09',preset='standard_business_brief',header='memo_masthead',
                 named_overrides={'CJK_font':'Microsoft YaHei','table_body_pt':9,'table_paragraph_after_pt':0,'caption_pt':9,'page_title_pt':16},pages=pages)
    text=json.dumps(payload,ensure_ascii=False,indent=2)
    if '导师' in text: raise ValueError('forbidden wording in objective document')
    (OUT/'document_content.json').write_text(text,encoding='utf8')

    # 完整技术报告共用Word内容的事实，另附审稿式边界与所有数值入口。
    md=['# 相似场景持续稳定阶段、人员组合与图像特征研究验证报告','',
        '2026-09-09。此报告对应本次全楼验证；旧uNb及h5窗口结果保留为历史基线。原始导出和正式Paper A方法合同未改。','',
        '## 主要结论与研究定位','',
        '**数据支持继续研究“不同图像的稳定阶段及其预测”，尚不支持把同一building看成具有一个统一停止人数。** 当前有同楼早/晚阶段和稳定多簇的实例；跨图预测中图像风险有可检验信号，但q主分析的改善方向仍受未知状态影响。','',
        '**按人独立采集、之后组合真实作答已经可行；稳定的人群类型尚未得到同等程度的支持。** 连续人员差异比固定两/三档复现更好。四人组合的结果可复算，但不能把它推广为某类20人收敛。','',
        '### 与研究交流记录的关系','',
        '导师书面确认的是：以人头为单位独立采集；ABCD表示人员类型；ABCD、AABC等比例在后处理中抽取组合，不为每种比例反复采集。书面记录还提出人数迭代的瓶颈/上限及不同场景下的含义。语音回忆中的building预测、风险与不确定性建模属于方向线索；研究问题表述和实验步骤在原文中仍需共同确定。不能把某个旧AI给出的分类或停止人数当成已经批准的最终RQ。','',
        '本轮把可检验问题落实为：RQ1：有限历史人数下，图像何时进入持续稳定阶段、参考质量边际变化如何；RQ2：同楼历史和图像特征对新图阶段有多少增量预测信息；RQ3：楼外形成的人员差异/档位是否复现、人员配比如何改变分歧与代表偏差。这是本轮操作化，不是替交流记录编造最终论文提纲。','']
    for page_no,page in enumerate(pages,1):
        md += [f'## {page_no}. {page["title"]}','']
        for block in page['blocks']:
            if block['type']=='p': md += [block['text'],'']
            elif block['type']=='table': md += [markdown_table(block['headers'],block['rows']),'']
            elif block['type']=='figure': md += [f'![{block["caption"]}]({Path(block["path"]).relative_to(OUT).as_posix()})','']
            elif block['type']=='link': md += [f'[{block["text"]}]({block["url"]})','']
    md += ['## 附录A：q与算法全部数值','',markdown_table(['人员口径','配置','起点可识别','观察内未达到','不确定'],qrows),'',
        '完整6q加OSPA及凝聚complete-link数值见 `stages/q_sensitivity_summary.csv`；上表省略.85/.80行仅为便读，不在机器表中排除。','',
        '## 附录B：主预测数值与比较边界','',markdown_table(['预测变量与来源','含W19/W26误差范围','排除W19/W26误差范围'],pred_rows),'',
        '这些线段为可行误差区间，不是给未知状态取中点后的单一MAE，也不是置信区间。配对改善共享同一未知目标；不同预测中重叠源的未知状态在区间计算中保守放宽，故区间可能较宽；来源身份与线性权重完全相同的预测，改善精确为0，不把恒等预测描述成效果不确定。','',
        '楼外全池均值具有比同楼小样本更多的历史图；此对照评价具体可用预测器，不能单独识别building因果效应。源图数量完全匹配的楼外随机/风险匹配对照保留于前一uNb结果包，和新持续阶段的结果不混算。不同模型的条件人数MAE必须检查共同覆盖，未把预测不到的晚阶段图删除后宣称普遍准确。','',
        '## 附录C：缺口与限制','',
        '- 低人数98图的长期标签尚不存在；61张N≤2只保留库存。低人数分支是可审核的前瞻预测，不是新增标注已经完成。','- 2条补点具有参照他人完整标注的派生来源。撤回补点已重算40个图×配置×终点组合；本轮未对全部预测模型和人员分类逐个重跑撤回补点，因此不声称所有效果已通过该敏感性。','- q/.10/80%和至少后续5人均为探索规则。q网格及达标率网格已经计算，当前未建立人工同/异簇标定集；未证明任何阈值唯一正确。','- 同批历史人员跨图共用；同楼留图检验新图预测，不是新人泛化。200排列与重复图划分非独立实验；条件building重抽样不重新抽人或重新训练模型。','- 参考未定403条不计质量最差；OSPA端点集不评价全部墙体连接、真实意图或语义。分歧大不等于低质量，稳定多簇不等于失败。','- 旧d_t完整参考池/可核实的等覆盖分数未在本轮输入中取得；旧risk有70图，其中高人数4图，不能冒名为d_t与新分数作公平优劣排名。','- 先前另一AI的第二份分析原始运行包仍未提供；未将其未经核验的41图/12楼等声明作为本轮输入。本轮人员41图/12楼数字来自此次独立实算，不是沿用该声明。','',
        '## 附录D：文件与复现','',
        '方法入口：[相似场景标注稳定性分析SOP](../../docs/thesis_main/相似场景标注稳定性分析SOP.md)。人员详细定义、成员、参考选择与全参数结果见[人员报告](workers/README_ZH.md)。','',
        '```powershell',
        '.venv/Scripts/python.exe tools/thesis_main/analysis/validate_building_stages_20260909.py features',
        '.venv/Scripts/python.exe tools/thesis_main/analysis/validate_building_stages_20260909.py replay',
        '.venv/Scripts/python.exe tools/thesis_main/analysis/validate_building_stages_20260909.py no-imputation',
        '.venv/Scripts/python.exe tools/thesis_main/analysis/predict_building_stages_20260909.py',
        '.venv/Scripts/python.exe tools/thesis_main/analysis/validate_worker_reuse_20260909.py',
        '.venv/Scripts/python.exe tools/thesis_main/analysis/build_research_validation_report_20260909.py',
        '```','',
        '图像特征运行额外安装PyYAML 6.0.2与torchvision 0.25.0至仓库已有虚拟环境，未改生产依赖。模型为torch 2.10.0+cpu；没有训练或替换checkpoint。Word用Codex捆绑Python及python-docx生成，渲染检查和最终测试结果见 `DELIVERY_CHECKS.json`。','',
        '字段合同：`stages/stage_curves.csv.gz`一行为图×人员口径×配置×H×阶段种类×k，三状态相加必须200；`lower/upper`由状态数直接计算。`image_stage_onsets.csv`在合法k≥2域提取possible/conservative/identified/status。`prediction/target_predictions.csv.gz`一行为目标×划分×模型，`count_error`缺失表示没有该条件人数评分而非0；`curve_error_*`与`gain_over_base_*`不是置信区间。`per_image_scores`先平均同一目标的出现，`per_building_scores`再按图平均，主汇总按楼平均。计数分母另列count_buildings/count_targets/count_appearances；总体曲线分母为buildings/images/target_appearances。`low_n_prospective_predictions.csv.gz`始终标记pending_long_horizon_annotations。所有image_id回连214图库存；workers字段详见其README。','']
    (OUT/'研究验证报告.md').write_text('\n'.join(md),encoding='utf8')
    print(f'content generated: {len(pages)} Word sections',flush=True)


def build_docx():
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    payload=json.loads((OUT/'document_content.json').read_text(encoding='utf8'))
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(.492)
    def style(name,size,color='20252B',before=0,after=6,bold=False):
        s=doc.styles[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.bold=bold; s.font.italic=False; s.font.color.rgb=RGBColor.from_string(color)
        s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
        pf=s.paragraph_format; pf.space_before=Pt(before); pf.space_after=Pt(after); pf.line_spacing=1.10
        pf.widow_control=True
        snap=OxmlElement('w:snapToGrid'); snap.set(qn('w:val'),'0'); s._element.get_or_add_pPr().append(snap)
        for border in list(s._element.get_or_add_pPr().findall(qn('w:pBdr'))):
            s._element.get_or_add_pPr().remove(border)
        return s
    style('Normal',11); style('Title',23,'20252B',after=8,bold=True); style('Subtitle',11,'596673',after=12)
    for name,size,color,before,after in [('Heading 1',16,'2E74B5',16,8),('Heading 2',13,'2E74B5',12,6),('Heading 3',12,'1F4D78',8,4)]:
        style(name,size,color,before,after,True).paragraph_format.keep_with_next=True
    style('Caption',9,'596673',4,6); style('Header',9,'596673',0,0); style('Footer',9,'596673',0,0)
    for name in ['Table Text','Table Header']:
        if name not in doc.styles: doc.styles.add_style(name,1)
        style(name,9,'20252B',0,0,name=='Table Header')
    sec.header.paragraphs[0].text=payload.get('header','相似场景标注稳定阶段 · 历史数据验证')
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(payload.get('date','2026-09-09')+'   |   ')
    for code in ('PAGE',):
        fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),code); footer._p.append(fld)
    def node(parent,name,attrs):
        e=OxmlElement('w:'+name)
        for k,v in attrs.items(): e.set(qn('w:'+k),str(v))
        parent.append(e); return e
    def add_table(block):
        rows=[block['headers'],*block['rows']]; n=len(rows[0]); table=doc.add_table(rows=len(rows),cols=n)
        table.autofit=False
        widths=[9360//n]*n; widths[-1]=9360-sum(widths[:-1])
        if n==2: widths=[2600,6760]
        if n==3: widths=[2350,2650,4360]
        if n==4: widths=[2700,2200,2200,2260]
        pr=table._tbl.tblPr
        for tag in ('tblW','tblInd','tblCellMar','tblBorders'):
            for child in list(pr.findall(qn('w:'+tag))): pr.remove(child)
        node(pr,'tblW',{'w':9360,'type':'dxa'}); node(pr,'tblInd',{'w':120,'type':'dxa'})
        mar=node(pr,'tblCellMar',{})
        for k,v in [('top',80),('bottom',80),('start',120),('end',120)]: node(mar,k,{'w':v,'type':'dxa'})
        borders=node(pr,'tblBorders',{})
        for k in ('top','left','bottom','right','insideH','insideV'): node(borders,k,{'val':'single','sz':4,'color':'D7DEE5'})
        grid=table._tbl.tblGrid
        for child in list(grid): grid.remove(child)
        for w in widths: node(grid,'gridCol',{'w':w})
        for r,values in enumerate(rows):
            trpr=table.rows[r]._tr.get_or_add_trPr(); node(trpr,'cantSplit',{})
            if r==0: node(trpr,'tblHeader',{})
            for j,text in enumerate(values):
                cell=table.cell(r,j); cell.width=Inches(widths[j]/1440)
                tcpr=cell._tc.get_or_add_tcPr(); w=tcpr.find(qn('w:tcW')); w.set(qn('w:w'),str(widths[j])); w.set(qn('w:type'),'dxa')
                cell.text=str(text); para=cell.paragraphs[0]; para.style=doc.styles['Table Header' if r==0 else 'Table Text']
                para.paragraph_format.keep_with_next=(r==0)
                if r==0: node(tcpr,'shd',{'fill':'F2F4F7'})
    for i,page in enumerate(payload['pages']):
        heading=doc.add_paragraph(page['title'],style='Title' if i==0 else 'Heading 1')
        if i: heading.paragraph_format.page_break_before=True
        if page.get('subtitle'): doc.add_paragraph(page['subtitle'],style='Subtitle')
        for b in page['blocks']:
            if b['type']=='p': doc.add_paragraph(b['text'])
            elif b['type']=='table': add_table(b)
            elif b['type']=='figure':
                para=doc.add_paragraph(); para.paragraph_format.keep_with_next=True; para.paragraph_format.space_after=Pt(0)
                if i == 0: para.paragraph_format.space_before=Pt(6)
                para.add_run().add_picture(b['path'],width=Inches(b.get('width',6.4)))
                doc.add_paragraph(b['caption'],style='Caption')
            elif b['type']=='link':
                para=doc.add_paragraph(); link=OxmlElement('w:hyperlink'); link.set(qn('r:id'),para.part.relate_to(b['url'],RT.HYPERLINK,is_external=True))
                run=node(link,'r',{}); pr=node(run,'rPr',{}); node(pr,'color',{'val':'2E74B5'}); node(run,'t',{}).text=b['text']; para._p.append(link)
    target=OUT/payload.get('filename','相似场景标注稳定阶段与人员组合_客观数据报告.docx')
    doc.core_properties.title=payload['pages'][0]['title']
    doc.core_properties.subject='历史数据与探索性验证'
    doc.core_properties.author='研究项目组'
    doc.save(target)
    # 结构审计：精确页面、样式和所有表格宽度；排除目标文档禁用字样。
    import zipfile,xml.etree.ElementTree as ET
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with zipfile.ZipFile(target) as z:
        xml=z.read('word/document.xml').decode('utf8')
        if '导师' in xml: raise ValueError('forbidden wording')
        tree=ET.fromstring(xml)
        for t in tree.findall('.//w:tbl',ns):
            ws=[int(c.get('{'+ns['w']+'}w')) for c in t.findall('./w:tblGrid/w:gridCol',ns)]
            assert sum(ws)==9360
            for row in t.findall('./w:tr',ns):
                assert [int(c.find('./w:tcPr/w:tcW',ns).get('{'+ns['w']+'}w')) for c in row.findall('./w:tc',ns)]==ws
    (OUT/'DOCX_STRUCTURE_QA.json').write_text(json.dumps(dict(status='passed',preset=payload['preset'],
        pages_planned=len(payload['pages']),tables=len(doc.tables),inline_images=len(doc.inline_shapes),
        table_width_dxa=9360,forbidden_word_absent=True,visual_qa='pending_render'),ensure_ascii=False,indent=2),encoding='utf8')
    print(target,flush=True)


if __name__=='__main__':
    build_docx() if '--docx' in sys.argv else build_content()
